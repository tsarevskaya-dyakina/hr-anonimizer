"""Независимые проверки чисел и дат на вымышленных данных."""
import base64
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
import io
import json
from pathlib import Path
import shutil
import subprocess
import unittest
import zipfile

import openpyxl
import xlsxwriter
import test_stage1 as stage1


def fixture(path, date1904=False):
    wb = xlsxwriter.Workbook(path, {'date_1904': date1904})
    ws = wb.add_worksheet('Числа')
    money = wb.add_format({'num_format': '#,##0.00', 'bold': True})
    date = wb.add_format({'num_format': 'dd.mm.yyyy hh:mm'})
    ws.write_row('A1', ['Оклад', 'Премия', 'Диапазон', 'Дата', 'Код', 'Расчёт'])
    for row, values in enumerate([[150111.75, 15000.25, 175123.75, datetime(2024, 2, 28, 12, 30), '00123'],
                                   [200111.50, 20000.50, 200000, datetime(2024, 2, 29, 12, 30), '00007'],
                                   [-2511.75, 0, -1, datetime(2024, 3, 1, 12, 30), '00000']], 1):
        for column, value in enumerate(values):
            if column in (0, 1):
                ws.write_number(row, column, value, money)
            elif column == 3:
                ws.write_datetime(row, column, value, date)
            else:
                ws.write(row, column, value)
    ws.write_formula('F2', '=A2+B2', money, 165112.0)
    ws.set_column('A:F', 22)
    ws.freeze_panes(1, 0)
    hidden = wb.add_worksheet('Скрыто')
    hidden.very_hidden()
    hidden.write_row('A1', ['Дата', 'Оклад'])
    hidden.write_datetime('A2', datetime(2024, 3, 2, 12, 30), date)
    hidden.write_number('B2', 101111.25, money)
    wb.close()


class Stage3(unittest.TestCase):
    setUpClass = classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass = classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp = stage1.Stage1.setUp
    tearDown = stage1.Stage1.tearDown

    def process(self, modes=None, scope='column', date1904=False):
        source = self.work / 'numbers.xlsx'
        fixture(source, date1904)
        modes = modes or ['scale', 'scale', 'range', 'date', 'keep', 'keep']
        payload = self.page.evaluate('''async ({data,modes,scope})=>{
          const S=SafeCycle,book=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);
          const settings=book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:
            Array.from({length:s.maxCol},(_,col)=>({mode:i?(col===0?'date':'scale'):modes[col],category:'Скрыто',rangeStep:50000}))}));
          const out=await S.Engine.anonymize(book,settings,'numbers.xlsx',false,undefined,undefined,{scaleScope:scope,crosscheck:false});
          const back=await S.Engine.restore(await out.blob.arrayBuffer(),JSON.parse(JSON.stringify(out.key)));
          async function b64(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}
          return {key:out.key,anon:await b64(out.blob),back:await b64(back.blob),report:{...back,blob:null},ranged:out.ranged};
        }''', {'data': base64.b64encode(source.read_bytes()).decode(), 'modes': modes, 'scope': scope})
        return source, payload

    def test_decimal_arithmetic_matches_python_decimal(self):
        values = ['1.005', '-1.005', '0', '150111.75', '1.23e-5', '-9.99e7', '9007199254740993.25']
        for inverse in (False, True):
            expected = []
            for raw in values:
                value = Decimal(raw) / Decimal('1.2345') if inverse else Decimal(raw) * Decimal('1.2345')
                expected.append(value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP))
            actual = self.page.evaluate('''({values,inverse})=>values.map(v=>SafeCycle.Numeric.scale(v,1.2345,2,inverse))''',
                                        {'values': values, 'inverse': inverse})
            self.assertEqual([Decimal(v) for v in actual], expected)

    def test_date_format_rules_and_time_only_exclusions(self):
        result = self.page.evaluate('formats=>formats.map(x=>SafeCycle.Numeric.formatInfo(x))',
                                    ['dd.mm.yyyy', 'yyyy-mm-dd hh:mm:ss', '[$-419]d mmmm yyyy', 'mmm-yy',
                                     '0 "days"', '0 \\d', '[Red]#,##0.00', '[h]:mm:ss', 'h:mm', 'mm:ss', '0.00%'])
        self.assertTrue(all(item['isDate'] for item in result[:4]))
        self.assertTrue(all(not item['isDate'] for item in result[4:]))
        self.assertTrue(result[7]['isTime'])
        self.assertEqual(result[-1]['decimals'], 4)

    def test_scale_common_and_column_factors_and_rounding(self):
        for scope in ('common', 'column'):
            source, payload = self.process(scope=scope)
            anonymous = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['anon'])))
            original = openpyxl.load_workbook(source)
            params = [p for p in payload['key']['transforms']['columns'] if p['mode'] == 'scale']
            self.assertEqual(len(params), 3)
            for param in params:
                self.assertGreaterEqual(param['factor'], .6)
                self.assertLessEqual(param['factor'], 1.4)
                self.assertNotEqual(param['factor'], 1)
                self.assertEqual(param['precision'], 2)
                for row in range(2, 5) if param['sheet'] == 'Числа' else [2]:
                    before = original[param['sheet']].cell(row, param['column'] + 1).value
                    after = anonymous[param['sheet']].cell(row, param['column'] + 1).value
                    expected = (Decimal(str(before)) * Decimal(str(param['factor']))).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
                    self.assertEqual(Decimal(str(after)), expected)
            if scope == 'common':
                self.assertEqual(len({p['factor'] for p in params}), 1)

    def test_exact_roundtrip_preserves_types_styles_fractional_dates_and_formula(self):
        source, payload = self.process()
        original = openpyxl.load_workbook(source)
        restored = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['back'])))
        anonymous = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['anon'])))
        for sheet in original:
            for row in sheet:
                for cell in row:
                    back = restored[sheet.title][cell.coordinate]
                    if sheet.title == 'Числа' and cell.column == 3 and cell.row > 1:
                        self.assertEqual(back.value, anonymous[sheet.title][cell.coordinate].value)
                    else:
                        self.assertEqual(back.value, cell.value)
                        self.assertEqual(back.data_type, cell.data_type)
                    self.assertEqual(back.number_format, cell.number_format)
                    self.assertEqual(back.style_id, cell.style_id)
        self.assertEqual(anonymous['Числа']['E2'].value, '00123')
        self.assertEqual(anonymous['Числа']['F2'].value, '=A2+B2')
        self.assertEqual(restored['Скрыто'].sheet_state, 'veryHidden')
        self.assertEqual(restored['Числа'].freeze_panes, 'A2')
        days = payload['key']['transforms']['dateShift']
        self.assertTrue(30 <= abs(days) <= 365)
        for sheet_name, cell in [('Числа', 'D2'), ('Числа', 'D3'), ('Скрыто', 'A2')]:
            self.assertEqual(anonymous[sheet_name][cell].value - original[sheet_name][cell].value, timedelta(days=days))

    def test_date1904_epoch_and_date_intervals(self):
        source, payload = self.process(date1904=True)
        original = openpyxl.load_workbook(source)
        anonymous = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['anon'])))
        restored = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['back'])))
        self.assertTrue(payload['key']['transforms']['date1904'])
        self.assertEqual(anonymous.epoch, original.epoch)
        self.assertEqual(anonymous['Числа']['D3'].value - anonymous['Числа']['D2'].value, timedelta(days=1))
        self.assertEqual(restored['Числа']['D2'].value, original['Числа']['D2'].value)

    def test_iso_dates_and_calendar_boundaries(self):
        result = self.page.evaluate('''()=>{
          const N=SafeCycle.Numeric,errors=[];
          for(const fn of [()=>N.shiftExcel('60',30,false),()=>N.shiftExcel('2958465',365,false),
            ()=>N.shiftExcel('0',-30,true),()=>N.shiftIso('2024-02-30',30)])try{fn();errors.push(false)}catch{errors.push(true)}
          return {errors,iso:N.shiftIso('2024-02-29T12:30:45.123+03:00',365),
            forward:N.shiftExcel('59',1,false),back:N.shiftExcel('61',-1,false)};
        }''')
        self.assertTrue(all(result['errors']))
        self.assertEqual(result['iso'], '2025-02-28T12:30:45.123+03:00')
        self.assertEqual(result['forward'], '61')
        self.assertEqual(result['back'], '59')

    def test_ranges_exclusive_boundary_and_no_original_in_key(self):
        result = self.page.evaluate('''()=>['175000','200000','-1','0','1.25','-0.25'].map((v,i)=>SafeCycle.Numeric.range(v,i<4?50000:.5))''')
        self.assertEqual(result, ['150–200 тыс.', '200–250 тыс.', '-50–0 тыс.', '0–50 тыс.', '1–1.5', '-0.5–0'])
        _, payload = self.process()
        self.assertNotIn('175123.75', json.dumps(payload['key']))
        self.assertFalse(any(c['ref'].startswith('C') and c['sheet'] == 'Числа' for c in payload['key']['cells']))
        self.assertEqual(payload['report']['irreversible'], 3)

    def test_only_ranges_are_not_reversed(self):
        result = self.page.evaluate('''async()=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer());
          const settings=book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:!i&&c===2?'range':'keep',rangeStep:50000}))}));
          const out=await S.Engine.anonymize(book,settings,'range.xlsx',false,undefined,undefined,{crosscheck:false}),back=await S.Engine.restore(await out.blob.arrayBuffer(),out.key);
          return {cells:out.key.cells,restored:back.restored,irreversible:back.irreversible,warnings:back.warnings};
        }''')
        self.assertEqual(result['cells'], [])
        self.assertEqual(result['restored'], 0)
        self.assertEqual(result['irreversible'], 3)
        self.assertTrue(result['warnings'])

    def test_text_codes_unformatted_dates_and_formulas_rejected(self):
        result = self.page.evaluate('''async()=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer()),errors=[];
          for(const [col,mode] of [[4,'scale'],[0,'date'],[3,'scale'],[4,'range']]){
            const settings=book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:!i&&c===col?mode:'keep',rangeStep:50000}))}));
            try{await S.Engine.anonymize(book,settings,'bad.xlsx');errors.push('')}catch(e){errors.push(e.message)}
          }return {errors,code:S.Xlsx.value(book.sheets[0].byRow.get(2).get(4),book.strings)};
        }''')
        self.assertTrue(all(result['errors']))
        self.assertEqual(result['code'], '00123')
        # Формулы проверяются для каждого нового режима.
        result = self.page.evaluate('''async()=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer()),errors=[];
          for(const mode of ['scale','range','date'])try{await S.Engine.anonymize(book,book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:!i&&c===3?mode:'keep',rangeStep:50000}))})),'formula.xlsx');errors.push('')}catch(e){errors.push(e.message)}return errors;
        }''')
        self.assertTrue(all('формула' in e for e in result))

    def test_edited_return_restores_reordered_columns_and_user_selected_new_calculation(self):
        _, payload = self.process(scope='common')
        source = next(p for p in payload['key']['transforms']['columns'] if p['sheet'] == 'Числа' and p['column'] == 0)
        edited = self.work / 'edited.xlsx'
        wb = xlsxwriter.Workbook(edited)
        ws = wb.add_worksheet('Ответ ИИ')
        money = wb.add_format({'num_format': '0.00'})
        date = wb.add_format({'num_format': 'yyyy-mm-dd hh:mm'})
        ws.write_row('A1', ['Новый расчёт', 'Оклад', 'Дата', 'Код'])
        scaled = lambda v: float((Decimal(str(v)) * Decimal(str(source['factor']))).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))
        ws.write_number('A2', scaled(1000), money)
        ws.write_number('B2', scaled(150111.75), money)
        ws.write_datetime('C2', datetime(2024, 2, 28, 12, 30) + timedelta(days=payload['key']['transforms']['dateShift']), date)
        ws.write('D2', '00123')
        ws.write_formula('B3', '=B2*2', money, scaled(300223.50))
        wb.close()
        result = self.page.evaluate('''async ({data,key,sourceId})=>{
          const S=SafeCycle,buffer=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,book=await S.Xlsx.load(buffer),candidates=S.Numeric.candidates(book,key),selected=candidates.filter(c=>c.selected);
          selected.push({sheet:'Ответ ИИ',column:0,sourceId});
          const back=await S.Engine.restore(buffer,key,undefined,undefined,{transforms:selected});
          const base64=await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(back.blob)});
          return {base64,candidates:candidates.map(c=>({column:c.column,selected:c.selected})),skipped:back.formulaSkipped};
        }''', {'data': base64.b64encode(edited.read_bytes()).decode(), 'key': payload['key'], 'sourceId': source['id']})
        restored = openpyxl.load_workbook(io.BytesIO(base64.b64decode(result['base64'])))
        self.assertFalse(result['candidates'][0]['selected'])
        self.assertEqual(restored['Ответ ИИ']['A2'].value, 1000)
        self.assertLessEqual(abs(restored['Ответ ИИ']['B2'].value - 150111.75), .02)
        self.assertEqual(restored['Ответ ИИ']['C2'].value, datetime(2024, 2, 28, 12, 30))
        self.assertEqual(restored['Ответ ИИ']['D2'].value, '00123')
        self.assertEqual(restored['Ответ ИИ']['B3'].value, '=B2*2')
        self.assertEqual(result['skipped'], 1)

    def test_invalid_key_parameters_and_missing_original_records_rejected(self):
        _, payload = self.process()
        result = self.page.evaluate('''key=>{
          const errors=[];for(const kind of ['factor','days','cells','count']){const copy=structuredClone(key),scale=copy.transforms.columns.find(p=>p.mode==='scale'),date=copy.transforms.columns.find(p=>p.mode==='date');
            if(kind==='factor')scale.factor=0;if(kind==='days')date.days=0;if(kind==='cells')copy.cells=[];if(kind==='count')scale.count=-1;
            try{SafeCycle.Key.validate(copy);errors.push(false)}catch{errors.push(true)}
          }return errors;
        }''', payload['key'])
        self.assertTrue(all(result))

    def test_ambiguous_headers_with_different_factors_require_selection(self):
        _, payload = self.process(scope='common')
        result = self.page.evaluate('''async key=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer());
          book.sheets[0].name='Ответ';const source=key.transforms.columns.filter(p=>p.mode==='scale'&&p.header==='Оклад');
          source[0].factor=.6;source[1].factor=1.4;
          return S.Numeric.candidates(book,key).filter(c=>c.sheet==='Ответ').map(c=>({header:c.header,selected:c.selected}));
        }''', payload['key'])
        self.assertFalse(next(c for c in result if c['header'] == 'Оклад')['selected'])

    def test_row_and_column_inherited_date_formats(self):
        source = self.work / 'inherit.xlsx'
        fixture(source)
        result = self.page.evaluate('''async data=>{
          const S=SafeCycle,book=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),sheet=book.sheets[0],first=sheet.byRow.get(2).get(3),style=first.getAttribute('s');
          first.removeAttribute('s');first.parentElement.setAttribute('s',style);first.parentElement.setAttribute('customFormat','true');
          sheet.byRow.get(3).get(3).removeAttribute('s');const col=S.Xml.create(sheet.doc,'col');col.setAttribute('min','4');col.setAttribute('max','4');col.setAttribute('style',style);S.Xml.all(sheet.doc,'cols')[0].append(col);
          const infos=[2,3].map(r=>S.Numeric.info(book,sheet,sheet.byRow.get(r).get(3)));
          const settings=book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:!i&&c===3?'date':'keep'}))}));
          const out=await S.Engine.anonymize(book,settings,'inherit.xlsx',false,undefined,undefined,{crosscheck:false});
          return {infos,count:out.transformed};
        }''', base64.b64encode(source.read_bytes()).decode())
        self.assertTrue(all(info['isDate'] for info in result['infos']))
        self.assertEqual(result['count'], 3)

    def test_iso_date_cells_roundtrip_and_epoch_change_rejected(self):
        result = self.page.evaluate('''async()=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer()),sheet=book.sheets[0],cell=sheet.byRow.get(2).get(3),iso='2024-02-29T12:30:45.123+03:00';S.Numeric.write(cell,iso,'d');
          const settings=book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:!i&&c===3?'date':'keep'}))}));
          const out=await S.Engine.anonymize(book,settings,'iso.xlsx',false,undefined,undefined,{crosscheck:false});
          const back=await S.Engine.restore(await out.blob.arrayBuffer(),out.key),restored=await S.Xlsx.load(await back.blob.arrayBuffer()),original=restored.sheets[0].byRow.get(2).get(3);
          const modified=await S.Xlsx.load(await out.blob.arrayBuffer()),archive=modified.archive.clone(),props=S.Xml.create(modified.workbook,'workbookPr');props.setAttribute('date1904','1');modified.workbook.documentElement.insertBefore(props,modified.workbook.documentElement.firstChild);archive.set('xl/workbook.xml',S.Xml.serialize(modified.workbook));
          let error;try{await S.Engine.restore(await (await archive.write()).arrayBuffer(),out.key)}catch(e){error=e.message}
          return {iso:S.Xlsx.value(original,restored.strings),type:original.getAttribute('t'),error};
        }''')
        self.assertEqual(result['iso'], '2024-02-29T12:30:45.123+03:00')
        self.assertEqual(result['type'], 'd')
        self.assertIn('другая система дат', result['error'])

    def test_stage2_key_and_text_are_still_supported(self):
        old = self.browser.new_page()
        try:
            old.goto(self.url.replace('v0.6.0', 'v0.2.0'))
            payload = old.evaluate("async()=>await SafeCycle.Engine.anonymizeText('Иванов Иван Иванович, demo@example.test.',{},false)")
        finally:
            old.close()
        result = self.page.evaluate('async ({text,key})=>await SafeCycle.Engine.restoreText(text,key)', payload)
        self.assertEqual(result['text'], 'Иванов Иван Иванович, demo@example.test.')
        self.assertTrue(result['sameHash'])

    def test_all_numeric_selfchecks_pass(self):
        result = self.page.evaluate('async()=>await SafeCycle.NumericSelfTest.run()')
        self.assertEqual(len(result), 4)
        self.assertTrue(all(t['pass'] for t in result), result)

    def test_numeric_demo_and_outputs_open_in_libreoffice(self):
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if not executable:
            self.skipTest('LibreOffice отсутствует')
        _, payload = self.process()
        demo64 = self.page.evaluate('''async()=>{const blob=await SafeCycle.Demo.createNumbers();return await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}''')
        destination = self.work / 'stage3-lo'
        destination.mkdir(exist_ok=True)
        files = []
        for name, data in [('stage3-anon.xlsx', payload['anon']), ('stage3-back.xlsx', payload['back']), ('stage3-demo.xlsx', demo64)]:
            path = self.work / name
            path.write_bytes(base64.b64decode(data))
            openpyxl.load_workbook(path)
            files.append(path)
        run = subprocess.run([executable, f'-env:UserInstallation={self.work.as_uri()}/stage3-lo-profile', '--headless',
                              '--convert-to', 'xlsx', '--outdir', str(destination), *map(str, files)],
                             capture_output=True, text=True, timeout=60)
        self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
        for path in files:
            converted = destination / path.name
            self.assertTrue(converted.exists(), run.stdout + run.stderr)
            openpyxl.load_workbook(converted)
        self.assertEqual(openpyxl.load_workbook(destination / files[1].name)['Числа']['D2'].value, datetime(2024, 2, 28, 12, 30))

    def test_ui_modes_range_step_and_common_scale(self):
        self.page.locator('#numeric-demo').click()
        self.page.locator('#config-card').wait_for(state='visible')
        self.assertEqual(self.page.get_by_label('Режим Сотрудники A', exact=True).input_value(), 'scale')
        self.assertEqual(self.page.get_by_label('Режим Сотрудники D', exact=True).input_value(), 'date')
        self.page.locator('#scale-scope').select_option('common')
        self.page.get_by_label('Режим Сотрудники C', exact=True).select_option('range')
        self.page.get_by_label('Шаг диапазона Сотрудники C', exact=True).fill('50000')
        self.page.locator('#process').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertIn('Диапазоны необратимы', self.page.locator('#warnings').inner_text())
        self.assertFalse(self.page.get_by_label('Шаг диапазона Сотрудники C', exact=True).is_disabled())
        self.assertTrue(self.page.get_by_label('Шаг диапазона Сотрудники A', exact=True).is_disabled())
        self.page.locator('#ack').check()
        self.page.locator('#key-encryption').uncheck()
        with self.page.expect_download() as download:
            self.page.locator('#download-key').click()
        key = self.work / 'ui-stage3-key.json'
        download.value.save_as(key)
        self.assertEqual(json.loads(key.read_text())['transforms']['scaleScope'], 'common')
        self.page.set_viewport_size({'width': 800, 'height': 1000})
        self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))

    def test_ui_edited_numeric_return_selects_new_column(self):
        _, payload = self.process(scope='common')
        key = self.work / 'return-key.json'
        key.write_text(json.dumps(payload['key']), encoding='utf-8')
        source = next(p for p in payload['key']['transforms']['columns'] if p['mode'] == 'scale')
        edited = self.work / 'numeric-answer.xlsx'
        wb = xlsxwriter.Workbook(edited)
        ws = wb.add_worksheet('Ответ')
        money = wb.add_format({'num_format': '0.00'})
        ws.write_row('A1', ['Оклад', 'Новый расчёт'])
        ws.write_number('A2', 100 * source['factor'], money)
        ws.write_number('B2', 200 * source['factor'], money)
        wb.close()
        self.page.locator('#restore-tab').click()
        self.page.locator('#key-file').set_input_files(key)
        self.page.locator('#return-file').set_input_files(edited)
        self.page.locator('#restore-options').wait_for(state='visible')
        self.assertTrue(self.page.get_by_label('Вернуть числа Ответ A', exact=True).is_checked())
        self.assertFalse(self.page.get_by_label('Вернуть числа Ответ B', exact=True).is_checked())
        self.page.get_by_label('Правило возврата Ответ B', exact=True).select_option(source['id'])
        self.assertTrue(self.page.get_by_label('Вернуть числа Ответ B', exact=True).is_checked())
        self.page.locator('#restore').click()
        self.page.locator('#download-restored').wait_for(state='visible')
        with self.page.expect_download() as download:
            self.page.locator('#download-restored').click()
        restored = self.work / 'ui-numeric-restored.xlsx'
        download.value.save_as(restored)
        self.assertEqual(openpyxl.load_workbook(restored)['Ответ']['B2'].value, 200)
        self.assertIn('Чисел и дат возвращено: 2', self.page.locator('#restore-report').inner_text())


if __name__ == '__main__':
    unittest.main(verbosity=2)
