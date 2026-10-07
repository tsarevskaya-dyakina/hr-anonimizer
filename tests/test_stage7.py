"""Итоговые сценарии: изменённые метки, учебный пакет, доступность, возврат и экспорт."""
import base64
import csv
import io
import json
import shutil
import subprocess
import unittest
import zipfile

import openpyxl
import xlsxwriter
from docx import Document
from playwright.sync_api import expect
import test_stage1 as stage1
import test_stage5 as stage5


class Stage7(unittest.TestCase):
    setUpClass=classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass=classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp=stage1.Stage1.setUp
    tearDown=stage1.Stage1.tearDown

    def text_result(self):
        self.page.locator('#source-text').click()
        self.page.locator('#raw-text').fill('Написать Иванову Ивану Ивановичу: ivan@example.test. Клиент Северный ветер.')
        self.page.locator('#process-text').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')

    def training_result(self):
        self.page.locator('#training-demo').click();self.page.locator('#busy').wait_for(state='hidden')
        expect(self.page.locator('#merge-list select')).to_have_count(50)
        self.page.locator('#merge-unique').click()
        self.page.locator('#process-batch').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')

    def test_label_variants_cases_leading_zeroes_and_word_boundaries(self):
        out=self.page.evaluate('''()=>{const e={label:'Сотрудник_0001',category:'Сотрудник',original:'Иванов Иван Иванович'},entries=new Map([[e.label,e]]);return SafeCycle.Labels.replace('[Сотруднику_0001]; СОТРУДНИКОМ_01; сотрудник_1; «Сотрудника_000001»; XСотрудник_1; Сотрудник_1X; _Сотрудник_1; Сотрудник_2',entries)}''')
        self.assertEqual(out['restored'],4);self.assertEqual(out['unknown'],2)
        self.assertIn('XСотрудник_1; Сотрудник_1X; _Сотрудник_1; Сотрудник_2',out['value'])
        self.assertEqual(out['value'].count(stage5.NAME),4)

    def test_ambiguous_label_forms_are_not_guessed(self):
        out=self.page.evaluate('''()=>{const entries=new Map([['Сотрудник_0001',{label:'Сотрудник_0001',category:'Сотрудник',original:'Первый'}],['Сотруднику_0001',{label:'Сотруднику_0001',category:'Сотруднику',original:'Второй'}]]);return SafeCycle.Labels.replace('СОТРУДНИКУ_1; Сотрудник_0001; Сотруднику_0001',entries)}''')
        self.assertEqual(out,{'value':'СОТРУДНИКУ_1; Первый; Второй','restored':2,'unknown':1})

    def test_original_compact_and_declined_labels_reserve_numbers(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,text='Иванов Иван Иванович, сотруднику_1 и СОТРУДНИК_2',out=await S.Engine.anonymizeText(text,{},false),back=await S.Engine.restoreText(out.text,out.key);return {labels:out.key.entries.map(e=>e.label),back:back.text,audit:out.audit}}''')
        self.assertIn('Сотрудник_0003',result['labels']);self.assertEqual(result['back'],'Иванов Иван Иванович, сотруднику_1 и СОТРУДНИК_2');self.assertFalse(result['audit']['blocked'])

    def test_text_answer_with_modified_labels_restores_original_values(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,out=await S.Engine.anonymizeText('Иванов Иван Иванович; ivan@example.test',{},false),text=out.text.replace(/Сотрудник_0+(\\d+)/g,'Сотруднику_$1').replace(/Email_0+(\\d+)/g,'email_$1');return S.Engine.restoreText(text+' [Сотрудник_9999]',out.key)}''')
        self.assertEqual(result['restored'],2);self.assertEqual(result['unknown'],1);self.assertIn(stage5.NAME,result['text']);self.assertIn('ivan@example.test',result['text']);self.assertIn('[Сотрудник_9999]',result['text'])

    def test_xlsx_edited_numeric_alias_keeps_number_type(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer()),cfg=b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Скрыто'}))})),out=await S.Engine.anonymize(b,cfg,'test.xlsx',false),anon=await S.Xlsx.load(await out.blob.arrayBuffer()),sheet=anon.sheets[0],cell=sheet.byRow.get(2).get(0),label=S.Xlsx.value(cell,anon.strings);S.Xlsx.setText(cell,'[скрытого_'+Number(label.split('_').at(-1))+']');anon.archive.set(sheet.path,S.Xml.serialize(sheet.doc));const back=await S.Engine.restore(await (await anon.archive.write()).arrayBuffer(),out.key),restored=await S.Xlsx.load(await back.blob.arrayBuffer()),c=restored.sheets[0].byRow.get(2).get(0);return {value:S.Xlsx.value(c,restored.strings),type:c.getAttribute('t')||'n',original:S.Xlsx.value(b.sheets[0].byRow.get(2).get(0),b.strings)}}''')
        self.assertEqual(result['type'],'n');self.assertEqual(result['value'],result['original'])

    def test_word_modified_label_split_across_runs(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Docx.load(await (await S.Docx.demo()).arrayBuffer()),out=await S.Docx.anonymize(b,'word.docx',{},false),anon=await S.Docx.load(await out.blob.arrayBuffer()),doc=anon.docs.get('word/document.xml'),p=S.Docx.all(doc,'p')[0],person=out.key.entries.find(e=>e.category==='Сотрудник'),label='[СОТРУДНИКУ_'+Number(person.label.split('_').at(-1))+']';p.replaceChildren();for(const piece of [label.slice(0,5),label.slice(5,10),label.slice(10)]){const r=doc.createElementNS(S.Docx.W,'w:r'),t=doc.createElementNS(S.Docx.W,'w:t');t.textContent=piece;r.append(t);p.append(r)}anon.archive.set('word/document.xml',S.Xml.serialize(doc));const back=await S.Engine.restore(await (await anon.archive.write()).arrayBuffer(),out.key),book=await S.Docx.load(await back.blob.arrayBuffer());return {text:S.Docx.paragraph(S.Docx.all(book.docs.get('word/document.xml'),'p')[0]).text,original:person.original,restored:back.restored}}''')
        self.assertEqual(result['text'],result['original']);self.assertGreater(result['restored'],0)

    def test_version_06_text_and_batch_keys_remain_supported(self):
        old=self.browser.new_page()
        try:
            old.goto(self.url.replace('v1.0.0','v0.6.0'))
            data=old.evaluate('''async()=>{const S=SafeCycle,text=await S.Engine.anonymizeText('Иванов Иван Иванович',{},false),files=[new File([await S.Demo.create()],'one.xlsx'),new File([await S.Demo.create()],'two.xlsx')],items=await S.Batch.load(files),settings={};for(const i of items)settings[i.id]=i.book.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:c===0?'alias':'keep',category:'Сотрудник'}))}));const batch=await S.Batch.process(items,settings,{},false),blob=batch.outputs[0].blob,base64=await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)});return {text:{text:text.text,key:text.key},batch:batch.key,data:base64}}''')
        finally:old.close()
        result=self.page.evaluate('''async data=>{const S=SafeCycle,text=await S.Engine.restoreText(data.text.text,data.text.key),back=await S.Batch.restoreOne(Uint8Array.from(atob(data.data),c=>c.charCodeAt(0)).buffer,data.batch,{id:'file-1',format:'xlsx'});return {text:text.text,same:back.sameHash}}''',data)
        self.assertEqual(result,{'text':stage5.NAME,'same':True})

    def test_partial_load_keeps_valid_file_and_lists_errors(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,items=await S.Batch.load([new File(['bad'],'broken.xlsx'),new File([await S.Demo.create()],'good.xlsx'),new File(['bad'],'scan.pdf')]);return {names:items.map(i=>i.name),errors:items.errors}}''')
        self.assertEqual(result['names'],['good.xlsx']);self.assertEqual([e['file'] for e in result['errors']],['broken.xlsx','scan.pdf'])
        good=self.work/'valid.xlsx';wb=openpyxl.Workbook();wb.active.append(['ФИО']);wb.active.append([stage5.NAME]);wb.save(good);bad=self.work/'broken.xlsx';bad.write_bytes(b'bad')
        self.page.locator('#file').set_input_files([str(bad),str(good)]);self.page.locator('#busy').wait_for(state='hidden');expect(self.page.locator('#load-errors')).to_be_visible();self.assertIn('broken.xlsx',self.page.locator('#load-errors').inner_text());self.assertIn('valid.xlsx',self.page.locator('#batch-files').inner_text())

    def test_partial_restore_keeps_valid_result_and_reports_corrupt_file(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer()),cfg=b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:c===0?'alias':'keep',category:'Сотрудник'}))})),out=await S.Engine.anonymize(b,cfg,'good.xlsx',false),back=await S.Batch.restoreFiles([new File(['bad'],'broken.xlsx'),new File([out.blob],'good_anon.xlsx')],out.key);return {names:back.outputs.map(o=>o.name),errors:back.errors,restored:back.restored}}''')
        self.assertEqual(result['names'],['good_anon.xlsx']);self.assertEqual(result['errors'][0]['file'],'broken.xlsx');self.assertGreater(result['restored'],0)

    def test_training_sources_have_fifty_people_chart_pivot_and_correct_formats(self):
        result=self.page.evaluate('''async()=>{async function enc(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}const result=[];for(const f of await SafeCycle.Training.files())result.push({name:f.name,data:await enc(f)});return result}''')
        byname={f['name']:base64.b64decode(f['data']) for f in result};wb=openpyxl.load_workbook(io.BytesIO(byname['Сотрудники.xlsx']));self.assertEqual(wb['Сотрудники'].max_row,51);self.assertEqual(wb['Оклады'].sheet_state,'hidden');self.assertEqual(wb['Сотрудники']['B2'].value,'00001');self.assertEqual(len(wb['Сотрудники']._charts),1);self.assertEqual(len(wb['Сводная']._pivots),1);self.assertEqual(wb['Сотрудники']['J2'].value,'=SUM(E2:E51)');self.assertIsNotNone(wb['Сотрудники']['A2'].comment)
        leaves=openpyxl.load_workbook(io.BytesIO(byname['Отпуска.xlsx']));self.assertEqual(leaves.active.max_row,51);self.assertEqual(leaves.active['A2'].value,'Иванов И.И.')
        doc=Document(io.BytesIO(byname['Служебная_записка.docx']));self.assertIn('9900000017','\n'.join(p.text for p in doc.paragraphs))
        rows=list(csv.reader(io.StringIO(byname['Сотрудники.csv'].decode('cp1251')),delimiter=';'));self.assertEqual(len(rows),51);self.assertEqual(rows[1][1],'00001');self.assertIn(';',rows[1][2])

    def test_training_ui_shared_people_colors_and_individual_download_gate(self):
        self.training_result();self.assertIn('Обработано файлов: 4',self.page.locator('#report').inner_text());expect(self.page.locator('#override-field')).to_be_hidden();self.assertGreater(self.page.locator('#preview .label-mark').count(),10);self.assertGreater(self.page.locator('#preview-legend .label-mark').count(),2)
        buttons=self.page.locator('#individual-downloads button');self.assertEqual(buttons.count(),4)
        for b in buttons.all():expect(b).to_be_disabled()
        self.page.locator('#ack').check()
        for b in buttons.all():expect(b).to_be_enabled()
        with self.page.expect_download() as d:buttons.first.click()
        self.assertEqual(d.value.suggested_filename,'Сотрудники_anon.xlsx');path=self.work/'training-anon.xlsx';d.value.save_as(path);wb=openpyxl.load_workbook(path);self.assertTrue(wb['Сотрудники']['A2'].value.startswith('Сотрудник_'));self.assertEqual(len(wb['Сотрудники']._charts),1);self.assertEqual(len(wb['Сводная']._pivots),1)

    def test_storage_bundle_is_password_protected_and_never_default_ai_download(self):
        self.text_result();expect(self.page.locator('#download-storage')).to_be_disabled();self.page.locator('#ack').check();self.page.locator('#key-password').fill('Учебный-пароль-2026!')
        with self.page.expect_download() as d:self.page.locator('#download-storage').click()
        self.assertEqual(d.value.suggested_filename,'Безопасный_цикл_для_хранения.zip');path=self.work/'storage.zip';d.value.save_as(path)
        with zipfile.ZipFile(path) as z:
            self.assertIn('Текст_anon.txt',z.namelist());name=next(n for n in z.namelist() if n.endswith('.json'));data=z.read(name);envelope=json.loads(data);self.assertEqual(envelope['schema'],'bezopasnyj-cikl-encrypted-key');self.assertNotIn(stage5.NAME,data.decode());self.assertIn('Никогда не передавайте',z.read('ПРОЧИТАЙТЕ.txt').decode())
        with self.page.expect_download() as d:self.page.locator('#download-file').click()
        self.assertTrue(d.value.suggested_filename.endswith('_anon.txt'));self.assertNotIn('key',d.value.suggested_filename)

    def test_storage_bundle_rejects_plaintext_and_short_password(self):
        self.text_result();self.page.locator('#ack').check();self.page.locator('#key-encryption').uncheck();self.page.locator('#download-storage').click();self.assertIn('включите защиту',self.page.locator('#status').inner_text());self.page.locator('#key-encryption').check();self.page.locator('#key-password').fill('short');self.page.locator('#download-storage').click();self.assertIn('не короче 8',self.page.locator('#status').inner_text())

    def test_ai_instruction_manual_copy_fallback_is_available(self):
        self.text_result();self.page.evaluate("Object.defineProperty(navigator,'clipboard',{value:undefined,configurable:true})");self.page.locator('#copy-ai').click();expect(self.page.locator('#ai-instruction-field')).to_be_visible();value=self.page.locator('#ai-instruction').input_value();self.assertIn('не склоняй',value);self.assertIn('не пытайся',value.lower());self.assertNotIn(stage5.NAME,value);self.assertIn('вручную',self.page.locator('#status').inner_text());self.assertEqual(self.page.evaluate('document.activeElement.id'),'ai-instruction')

    def test_tabs_keyboard_offline_help_and_accessible_fields(self):
        self.page.locator('#anon-tab').focus();self.page.keyboard.press('ArrowRight');expect(self.page.locator('#restore-tab')).to_be_focused();expect(self.page.locator('#restore-panel')).to_be_visible();self.page.keyboard.press('Home');expect(self.page.locator('#anon-tab')).to_be_focused()
        self.page.locator('#offline-link').click();self.assertTrue(self.page.locator('#security-details').evaluate('n=>n.open'));self.assertEqual(self.page.evaluate('document.activeElement.id'),'offline-help')
        self.page.locator('#demo').click();self.page.locator('#busy').wait_for(state='hidden')
        missing=self.page.evaluate("[...document.querySelectorAll('input,select,textarea,progress')].filter(n=>n.getClientRects().length&&!n.disabled&&!n.labels?.length&&!n.getAttribute('aria-label')&&!n.getAttribute('aria-labelledby')).map(n=>n.id||n.outerHTML.slice(0,120))")
        self.assertEqual(missing,[])

    def test_preview_unchanged_toggle_and_categories_are_not_color_only(self):
        self.page.locator('#demo').click();self.page.locator('#busy').wait_for(state='hidden');self.page.locator('#process').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        rows=self.page.locator('#preview [data-unchanged]');self.assertGreater(rows.count(),0);self.assertTrue(all(not r.is_visible() for r in rows.all()));self.page.locator('#preview-show-all').check();self.assertTrue(all(r.is_visible() for r in rows.all()));marks=self.page.locator('#preview .label-mark');self.assertGreater(marks.count(),0);self.assertTrue(all(m.get_attribute('aria-label') and m.get_attribute('data-category') for m in marks.all()))

    def test_light_dark_contrast_and_narrow_layout(self):
        self.text_result()
        for scheme in ['light','dark']:
            self.page.emulate_media(color_scheme=scheme,reduced_motion='reduce')
            ratios=self.page.evaluate('''()=>{function rgb(s){return s.match(/[\\d.]+/g).slice(0,3).map(Number)}function luminance(c){return c.map(v=>{v/=255;return v<=.04045?v/12.92:((v+.055)/1.055)**2.4}).reduce((s,v,i)=>s+v*[.2126,.7152,.0722][i],0)}return [...document.querySelectorAll('#preview .label-mark,#preview-legend .label-mark')].map(n=>{const c=getComputedStyle(n),a=luminance(rgb(c.color)),b=luminance(rgb(c.backgroundColor));return (Math.max(a,b)+.05)/(Math.min(a,b)+.05)})}''')
            self.assertTrue(ratios);self.assertTrue(all(r>=4.5 for r in ratios),(scheme,ratios))
            for width in [1280,800,390]:
                self.page.set_viewport_size({'width':width,'height':1000});self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'),(scheme,width))

    def test_reset_clears_new_exports_prompt_and_password_from_return_tab(self):
        self.text_result();self.page.locator('#copy-ai').click();self.page.locator('#key-password').fill('Учебный-пароль');self.page.locator('#restore-tab').click();self.page.locator('#restore-password').fill('Ещё-пароль');self.page.locator('#reset').click();expect(self.page.locator('#anon-panel')).to_be_visible();expect(self.page.locator('#result-card')).to_be_hidden();self.assertEqual(self.page.locator('#key-password').input_value(),'');self.assertEqual(self.page.locator('#restore-password').input_value(),'');self.assertEqual(self.page.locator('#ai-instruction').input_value(),'');self.assertEqual(self.page.locator('#preview-legend').inner_text(),'');expect(self.page.locator('#download-storage')).to_be_disabled()

    def test_final_selfchecks_pass(self):
        results=self.page.evaluate('async()=>await SafeCycle.FinalSelfTest.run()');self.assertEqual(len(results),4);self.assertTrue(all(r['pass'] for r in results),results)

    def test_return_input_and_numeric_settings_invalidate_previous_download(self):
        data=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer()),cfg=b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:s.name==='Сотрудники'&&i===0?'scale':'keep',category:'Скрыто'}))})),out=await S.Engine.anonymize(b,cfg,'numbers.xlsx',false),a=await S.Xlsx.load(await out.blob.arrayBuffer());S.Xlsx.setText(a.sheets[0].byRow.get(1).get(0),'Новый расчёт');a.archive.set(a.sheets[0].path,S.Xml.serialize(a.sheets[0].doc));const blob=await a.archive.write(),encoded=await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)});return {key:out.key,data:encoded}}''')
        key=self.work/'numbers-key.json';key.write_text(json.dumps(data['key'],ensure_ascii=False));answer=self.work/'answer.xlsx';answer.write_bytes(base64.b64decode(data['data']));bad=self.work/'bad.xlsx';bad.write_bytes(b'broken')
        self.page.locator('#restore-tab').click();self.page.locator('#key-file').set_input_files(str(key));self.page.locator('#return-file').set_input_files(str(answer));self.page.locator('#restore-options').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden');flag=self.page.get_by_label('Вернуть числа Сотрудники A',exact=True);flag.check();self.page.locator('#restore').click();self.page.locator('#busy').wait_for(state='hidden');expect(self.page.locator('#download-restored')).to_be_visible();flag.uncheck();expect(self.page.locator('#download-restored')).to_be_hidden();expect(self.page.locator('#restore-report')).to_be_hidden();flag.check();self.page.locator('#restore').click();self.page.locator('#busy').wait_for(state='hidden');expect(self.page.locator('#download-restored')).to_be_visible();self.page.locator('#return-file').set_input_files(str(bad));expect(self.page.locator('#download-restored')).to_be_hidden();self.page.locator('#busy').wait_for(state='hidden')

    def test_word_only_package_hides_inapplicable_column_controls(self):
        source=self.work/'one.docx';source.write_bytes(base64.b64decode(self.page.evaluate('''async()=>{const b=await SafeCycle.Docx.demo();return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(b)})}''')));second=self.work/'two.docx';second.write_bytes(source.read_bytes());self.page.locator('#file').set_input_files([str(source),str(second)]);self.page.locator('#busy').wait_for(state='hidden');expect(self.page.locator('#config-card')).to_be_hidden();expect(self.page.locator('#apply-matching')).to_be_hidden();expect(self.page.locator('#word-options')).to_be_visible();self.page.locator('#process-batch').click();self.page.locator('#busy').wait_for(state='hidden');expect(self.page.locator('#result-card')).to_be_visible();self.assertIn('Обработано файлов: 2',self.page.locator('#report').inner_text())

    def test_training_outputs_and_restored_files_open_in_libreoffice(self):
        data=self.page.evaluate('''async()=>{const S=SafeCycle,items=await S.Batch.load(await S.Training.files()),{settings,options}=await S.Training.plan(items),out=await S.Batch.process(items,settings,options,false),result=[];async function encode(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}for(const o of out.outputs.filter(o=>o.format!=='csv')){result.push({name:S.Batch.name(o.name,'_anon',o.format),data:await encode(o.blob)});const back=await S.Batch.restoreOne(await o.blob.arrayBuffer(),out.key,{id:o.id,format:o.format});result.push({name:S.Batch.name(o.name,'_restored',o.format),data:await encode(back.blob)})}return result}''')
        files=[]
        for item in data:
            p=self.work/item['name'];p.write_bytes(base64.b64decode(item['data']));files.append(p)
        dest=self.work/'training-lo';dest.mkdir(exist_ok=True);exe=shutil.which('soffice');self.assertIsNotNone(exe)
        for extension in ['xlsx','docx']:
            selected=[p for p in files if p.suffix=='.'+extension];run=subprocess.run([exe,f'-env:UserInstallation={self.work.as_uri()}/training-profile','--headless','--convert-to',extension,'--outdir',str(dest),*map(str,selected)],capture_output=True,text=True,timeout=60);self.assertEqual(run.returncode,0,run.stdout+run.stderr)
            for p in selected:self.assertTrue((dest/p.name).exists(),run.stdout+run.stderr)
        wb=openpyxl.load_workbook(dest/'Сотрудники_anon.xlsx');self.assertEqual(wb['Сотрудники'].max_row,51);self.assertEqual(len(wb['Сотрудники']._charts),1);self.assertEqual(len(wb['Сводная']._pivots),1);self.assertEqual(wb['Оклады'].sheet_state,'hidden')

    def test_converted_single_macro_book_returns_actual_xlsx_format(self):
        out=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer()),cfg=b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Сотрудник'}))}));b.format='xlsm';const anon=await S.Engine.anonymize(b,cfg,'source.xlsm',false),back=await S.Engine.restore(await anon.blob.arrayBuffer(),anon.key);return {macro:back.isXlsm,restored:back.restored}}''')
        self.assertFalse(out['macro']);self.assertGreater(out['restored'],0)

    def test_real_xlsx_100000_rows_exact_return_and_responsiveness(self):
        path=self.work/'large-final.xlsx';wb=xlsxwriter.Workbook(path,{'constant_memory':True});ws=wb.add_worksheet('Сотрудники');ws.write_row(0,0,['ФИО','Код'])
        for i in range(100000):ws.write_row(i+1,0,[stage5.NAME,str(i).zfill(6)])
        wb.close();payload=base64.b64encode(path.read_bytes()).decode()
        out=self.page.evaluate('''async data=>{const S=SafeCycle,bytes=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer;let last=performance.now(),maxGap=0,beats=0,b;const interval=setInterval(()=>{const now=performance.now();maxGap=Math.max(maxGap,now-last);last=now;beats++},20),started=performance.now();try{b=await S.Xlsx.load(bytes);const cfg=[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'alias',category:'Сотрудник'},{mode:'keep',category:'Скрыто'}]}],anon=await S.Engine.anonymize(b,cfg,'large.xlsx',false);b=null;const back=await S.Engine.restore(await anon.blob.arrayBuffer(),anon.key),book=await S.Xlsx.load(await back.blob.arrayBuffer()),ms=performance.now()-started;clearInterval(interval);let correct=true;for(let r=2;r<=100001;r++){const row=book.sheets[0].byRow.get(r);if(S.Xlsx.value(row.get(0),book.strings)!=='Иванов Иван Иванович'||S.Xlsx.value(row.get(1),book.strings)!==String(r-2).padStart(6,'0')){correct=false;break}}return {rows:book.sheets[0].byRow.size,correct,audit:anon.audit.total,restored:back.restored,ms,maxGap,beats}}finally{clearInterval(interval)}}''',payload)
        self.assertEqual(out['rows'],100001);self.assertTrue(out['correct']);self.assertEqual(out['audit'],0);self.assertEqual(out['restored'],100000);self.assertGreater(out['beats'],30);self.assertLess(out['maxGap'],1500,out);print('XLSX 100000 exact:',json.dumps(out),flush=True)
