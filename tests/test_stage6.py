"""Пакет, CSV, криптография, ключ XLSX, Web Worker. Только вымышленные данные."""
import base64
import csv
from datetime import datetime, timedelta
import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import time
import unittest
import zipfile

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from docx import Document
import openpyxl
from playwright.sync_api import expect
import test_stage1 as stage1
import test_stage5 as stage5

NAME=stage5.NAME
EMAIL=stage5.EMAIL
PASSWORD='Учебный-пароль-2026!'


def workbook(path, short=False):
    wb=openpyxl.Workbook();s=wb.active;s.title='Сотрудники'
    s.append(['ФИО','Код','Оклад','Дата','Комментарий'])
    s.append(['Иванов И.И.' if short else NAME,'00123',150000.50,datetime(2026,10,7),'Передать Иванову Ивану Ивановичу; '+EMAIL])
    s['D2'].number_format='dd.mm.yyyy';s['C2'].number_format='#,##0.00';s['A2'].font=openpyxl.styles.Font(bold=True)
    wb.save(path)


class Stage6(unittest.TestCase):
    setUpClass=classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass=classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp=stage1.Stage1.setUp
    tearDown=stage1.Stage1.tearDown

    def csv_process(self, payload, modes=None, options=None, restore=True):
        result=self.page.evaluate('''async ({data,modes,options,restore})=>{const S=SafeCycle,b=await S.Csv.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),settings=[{path:b.sheets[0].path,headerRow:1,columns:Array.from({length:b.sheets[0].maxCol},(_,i)=>modes?.[i]||{mode:'scan',category:'Скрыто'})}],out=await S.Csv.anonymize(b,settings,'test.csv',false,null,null,options||{}),back=restore?await S.Engine.restore(await out.blob.arrayBuffer(),out.key):null;async function enc(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}return {key:out.key,audit:out.audit,anon:await enc(out.blob),back:back?await enc(back.blob):null,dialect:b.csv.dialect}}''',{'data':base64.b64encode(payload).decode(),'modes':modes,'options':options,'restore':restore})
        return result

    def bundle(self, modes=None, options=None, include_csv=True):
        paths=[self.work/'staff.xlsx',self.work/'leaves.xlsx',self.work/'memo.docx']
        workbook(paths[0]);workbook(paths[1],short=True);stage5.fixture(paths[2])
        if include_csv:
            path=self.work/'data.csv';path.write_bytes(('ФИО;Код;Оклад;Дата;Комментарий\r\n'+NAME+';00123;150000,50;2026-10-07;'+EMAIL+'\r\n').encode('cp1251'));paths.append(path)
        payload=[{'name':p.name,'data':base64.b64encode(p.read_bytes()).decode()} for p in paths]
        result=self.page.evaluate('''async ({payload,modes,options})=>{const S=SafeCycle,files=payload.map(p=>new File([Uint8Array.from(atob(p.data),c=>c.charCodeAt(0))],p.name)),items=await S.Batch.load(files),settings={};for(const item of items)if(item.format!=='docx')settings[item.id]=item.book.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>modes?.[i]||{mode:i===0?'alias':'scan',category:i===0?'Сотрудник':'Скрыто'})}));const texts=items.flatMap(i=>i.format==='docx'?S.Docx.sourceTexts(i.book):i.book.sheets.flatMap(s=>s.cells.map(c=>S.Xlsx.value(c,i.book.strings)))),catalog=S.Detectors.catalog(texts),merges=catalog.suggestions.filter(s=>s.candidates.length===1).map(s=>({shortId:s.shortId,fullId:s.candidates[0].id})),out=await S.Batch.process(items,settings,{merges,...options},false);async function enc(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}const outputs=[];for(const o of out.outputs){const b=await S.Batch.restoreOne(await o.blob.arrayBuffer(),out.key,{format:o.format,fileName:o.name});outputs.push({name:o.name,format:o.format,anon:await enc(o.blob),back:await enc(b.blob),sameHash:b.sameHash,audit:o.audit,sourceId:b.sourceId})}return {key:out.key,audit:out.audit,outputs,zip:await enc(await S.Batch.zip(out.outputs))}}''',{'payload':payload,'modes':modes,'options':options or {}})
        for output in result['outputs']:
            (self.work/('anon-'+output['name'])).write_bytes(base64.b64decode(output['anon']))
            (self.work/('back-'+output['name'])).write_bytes(base64.b64decode(output['back']))
        return result,paths

    def test_csv_utf8_bom_quoted_semicolon_multiline_and_exact_values(self):
        text='ФИО;Код;Комментарий\r\n"'+NAME+'";"00123";"Первая; строка\r\nОт '+NAME+' ""одобрено"""\r\n'
        result=self.csv_process(b'\xef\xbb\xbf'+text.encode(),[{'mode':'alias','category':'Сотрудник'},{'mode':'keep','category':'Скрыто'},{'mode':'scan','category':'Скрыто'}])
        self.assertEqual(result['dialect']['encoding'],'utf-8');self.assertTrue(result['dialect']['bom']);self.assertEqual(result['dialect']['separator'],';')
        anon=list(csv.reader(io.StringIO(base64.b64decode(result['anon']).decode('utf-8-sig')),delimiter=';'))
        back=list(csv.reader(io.StringIO(base64.b64decode(result['back']).decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(back,list(csv.reader(io.StringIO(text),delimiter=';')));self.assertEqual(anon[1][1],'00123');self.assertNotIn(NAME,anon[1][2]);self.assertFalse(result['audit']['blocked'])

    def test_csv_windows1251_comma_tab_and_separator_detection(self):
        for encoding,delimiter in [('cp1251',';'),('utf-8',','),('utf-8','\t')]:
            with self.subTest(encoding=encoding,delimiter=delimiter):
                text=delimiter.join(['ФИО','Код'])+'\n'+delimiter.join([NAME,'00123'])+'\n'
                result=self.csv_process(text.encode(encoding),[{'mode':'alias','category':'Сотрудник'},{'mode':'keep','category':'Скрыто'}]);self.assertEqual(result['dialect']['separator'],delimiter);self.assertEqual(result['dialect']['encoding'],'windows-1251' if encoding=='cp1251' else 'utf-8')
                self.assertEqual(base64.b64decode(result['back']).decode(encoding),text)

    def test_csv_numeric_and_iso_dates_preserve_original_lexemes(self):
        payload=('ФИО;Код;Сумма;Дата\n'+NAME+';00123;150000,50;2026-10-07\n').encode()
        result=self.csv_process(payload,[{'mode':'alias','category':'Сотрудник'},{'mode':'keep','category':'Скрыто'},{'mode':'scale','category':'Сумма'},{'mode':'date','category':'Дата'}]);self.assertEqual(base64.b64decode(result['back']),payload)
        anon=list(csv.reader(io.StringIO(base64.b64decode(result['anon']).decode()),delimiter=';'))
        self.assertNotEqual(float(anon[1][2]),150000.5);self.assertNotEqual(anon[1][3],'2026-10-07');self.assertEqual(anon[1][1],'00123')

    def test_csv_range_and_delete_never_enter_key(self):
        payload=('Секрет;Оклад;Комментарий\nprivate@example.test;175123;private@example.test и 175123\n').encode()
        result=self.csv_process(payload,[{'mode':'delete','category':'Email'},{'mode':'range','category':'Сумма','rangeStep':50000},{'mode':'scan','category':'Скрыто'}]);raw=json.dumps(result['key'],ensure_ascii=False)
        self.assertNotIn('private@example.test',raw);self.assertNotIn('175123',raw)
        self.assertNotIn('private@example.test',base64.b64decode(result['back']).decode());self.assertNotIn('175123',base64.b64decode(result['back']).decode())

    def test_csv_formula_risk_is_blocked_and_valid_negative_number_is_not(self):
        for value,blocked in [('=1+1',True),('@SUM(1)',True),('+cmd',True),('-100',False)]:
            with self.subTest(value=value):
                result=self.csv_process(('Значение\n'+value+'\n').encode(),[{'mode':'keep','category':'Скрыто'}],restore=False)
                self.assertEqual(result['audit']['blocked'],blocked,result['audit'])

    def test_csv_invalid_quotes_encoding_and_leading_zero_scale_rejected(self):
        errors=self.page.evaluate('''async()=>{const S=SafeCycle,errors=[];for(const text of ['ФИО;Код\\n"Иван;123','a;b\\n"x"z;1','a;b\\nfoo\\0;1'])try{await S.Csv.load(new TextEncoder().encode(text).buffer);errors.push(null)}catch(e){errors.push(e.message)}return errors}''')
        self.assertTrue(all(errors),errors)
        with self.assertRaises(Exception):self.csv_process(b'Code\n00123\n',[{'mode':'scale','category':'Сумма'}])

    def test_batch_same_people_across_excel_word_csv_and_initials(self):
        result,_=self.bundle()
        self.assertEqual(len(result['key']['files']),4);self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        people=[e for e in result['key']['entries'] if e['category']=='Сотрудник' and e['canonical']==NAME];self.assertEqual(len(people),1)
        label=people[0]['label'];self.assertEqual(openpyxl.load_workbook(self.work/'anon-staff.xlsx').active['A2'].value,label);self.assertEqual(openpyxl.load_workbook(self.work/'anon-leaves.xlsx').active['A2'].value,label)
        self.assertIn('['+label+']','\n'.join(p.text for p in Document(self.work/'anon-memo.docx').paragraphs))
        rows=list(csv.reader(io.StringIO((self.work/'anon-data.csv').read_bytes().decode('cp1251')),delimiter=';'));self.assertEqual(rows[1][0],label)
        self.assertEqual(openpyxl.load_workbook(self.work/'back-leaves.xlsx').active['A2'].value,'Иванов И.И.');self.assertTrue(all(o['sameHash'] for o in result['outputs']))
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['zip']))) as z:self.assertEqual(len(z.namelist()),4);self.assertTrue(all('_anon.' in n for n in z.namelist()))
        self.assertEqual({p['file'] for p in people[0]['locations']},{'staff.xlsx','leaves.xlsx','memo.docx','data.csv'})

    def test_batch_common_factor_and_one_date_shift_across_formats(self):
        modes=[{'mode':'alias','category':'Сотрудник'},{'mode':'keep','category':'Скрыто'},{'mode':'scale','category':'Сумма'},{'mode':'date','category':'Дата'},{'mode':'scan','category':'Скрыто'}]
        result,_=self.bundle(modes,{'scaleScope':'common'});params=[p for f in result['key']['files'] for p in f['key'].get('transforms',{}).get('columns',[])]
        self.assertEqual(len({p['factor'] for p in params if p['mode']=='scale'}),1);self.assertEqual(len({p['days'] for p in params if p['mode']=='date'}),1)
        days=next(p['days'] for p in params if p['mode']=='date');wb=openpyxl.load_workbook(self.work/'anon-staff.xlsx');self.assertEqual(wb.active['D2'].value,datetime(2026,10,7)+timedelta(days=days))
        self.assertEqual(openpyxl.load_workbook(self.work/'back-staff.xlsx').active['C2'].value,150000.5)
        self.assertIn('150000,50',(self.work/'back-data.csv').read_bytes().decode('cp1251'))

    def test_key_aes_gcm_pbkdf2_verified_independently(self):
        result,_=self.bundle(include_csv=False);key=result['key']
        encrypted=self.page.evaluate('''async ({key,password})=>{const blob=await SafeCycle.ProtectedKey.save(key,password);return JSON.parse(await blob.text())}''',{'key':key,'password':PASSWORD})
        self.assertGreaterEqual(encrypted['iterations'],310000);self.assertEqual(encrypted['algorithm'],'AES-256-GCM')
        salt=base64.b64decode(encrypted['salt']);iv=base64.b64decode(encrypted['iv']);self.assertEqual(len(salt),16);self.assertEqual(len(iv),12)
        derived=hashlib.pbkdf2_hmac('sha256',PASSWORD.encode(),salt,encrypted['iterations'],32)
        aad='|'.join(str(encrypted[k]) for k in ['schema','version','algorithm','kdf','iterations','format']).encode()
        decoded=json.loads(AESGCM(derived).decrypt(iv,base64.b64decode(encrypted['data']),aad));self.assertEqual(decoded,key)
        self.assertNotIn(NAME,json.dumps(encrypted,ensure_ascii=False));self.assertNotIn('entries',encrypted)
        # A separately produced ciphertext opens in the browser too.
        independent=dict(encrypted);independent['data']=base64.b64encode(AESGCM(derived).encrypt(iv,json.dumps(key,ensure_ascii=False).encode(),aad)).decode()
        opened=self.page.evaluate('''async ({e,password})=>await SafeCycle.ProtectedKey.open(new TextEncoder().encode(JSON.stringify(e)).buffer,password)''',{'e':independent,'password':PASSWORD});self.assertEqual(opened,key)

    def test_key_wrong_password_tampering_salts_and_iv(self):
        out=self.page.evaluate('''async()=>{const S=SafeCycle,key=(await S.Engine.anonymizeText('Передать '+ 'Иванов Иван Иванович')).key,a=JSON.parse(await (await S.ProtectedKey.save(key,'Password-123')).text()),b=JSON.parse(await (await S.ProtectedKey.save(key,'Password-123')).text());const errors=[];for(const [e,password] of [[a,'Wrong-123'],[{...a,format:'xlsx'},'Password-123'],[{...a,data:a.data.slice(0,-4)+'AAAA'},'Password-123'],[{...a,iterations:1},'Password-123']])try{await S.ProtectedKey.open(new TextEncoder().encode(JSON.stringify(e)).buffer,password);errors.push(null)}catch(e){errors.push(e.message)}return {a,b,errors}}''')
        self.assertNotEqual(out['a']['salt'],out['b']['salt']);self.assertNotEqual(out['a']['iv'],out['b']['iv']);self.assertTrue(all(out['errors']),out['errors'])

    def test_human_xlsx_key_and_encrypted_xlsx_roundtrip(self):
        result,_=self.bundle()
        data=self.page.evaluate('''async ({key,password})=>{const S=SafeCycle,a=await S.ProtectedKey.save(key,null,'xlsx'),b=await S.ProtectedKey.save(key,password,'xlsx'),opened=await S.ProtectedKey.open(await b.arrayBuffer(),password);return {plain:await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(a)}),opened}}''',{'key':result['key'],'password':PASSWORD})
        path=self.work/'human-key.xlsx';path.write_bytes(base64.b64decode(data['plain']));wb=openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames[:2],['Соответствия','Параметры']);self.assertEqual(wb['Ключ программы'].sheet_state,'veryHidden');self.assertIn(NAME,[row[1].value for row in wb['Соответствия'].iter_rows(min_row=2)]);self.assertEqual(data['opened'],result['key'])
        read=self.page.evaluate('''async data=>await SafeCycle.ProtectedKey.open(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer)''',data['plain']);self.assertEqual(read,result['key'])

    def test_batch_duplicate_output_names_and_ambiguous_edited_file(self):
        result,_=self.bundle(include_csv=False)
        error=self.page.evaluate('''async ({key,data})=>{const S=SafeCycle,b=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),a=b.archive.clone();a.set('docProps/extra.xml','<extra/>');try{await S.Batch.restoreOne(await (await a.write()).arrayBuffer(),key,{format:'xlsx',fileName:'renamed.xlsx'});return null}catch(e){return e.message}}''',{'key':result['key'],'data':result['outputs'][0]['anon']})
        self.assertIn('выберите',error)
        data=self.page.evaluate('''async()=>{const S=SafeCycle,outputs=[{name:'same.xlsx',format:'xlsx',blob:new Blob(['a'])},{name:'same.xlsx',format:'xlsx',blob:new Blob(['b'])}],z=await S.Zip.Archive.read(await (await S.Batch.zip(outputs)).arrayBuffer());return [...z.entries.keys()]}''');self.assertEqual(data,['same_anon.xlsx','same_anon_2.xlsx'])

    def test_batch_key_plain_text_return_uses_shared_correspondences(self):
        result,_=self.bundle(include_csv=False);label=next(e['label'] for e in result['key']['entries'] if e['canonical']==NAME)
        text=self.page.evaluate('''async ({key,text})=>await SafeCycle.Engine.restoreText(text,key)''',{'key':result['key'],'text':'Ответ: ['+label+']'});self.assertEqual(text['text'],'Ответ: '+NAME)

    def test_cancel_worker_terminates_and_next_job_succeeds(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,ctrl=new AbortController(),p=S.Background.run('plan',{records:Array.from({length:100000},(_,i)=>({id:String(i),sheet:'Тест',ref:'A'+(i+1),column:0,value:'Иванов Иван Иванович',mode:'scan',category:'Сотрудник',type:'inlineStr',isText:true})),options:{},shuffle:false},ctrl.signal);setTimeout(()=>ctrl.abort(),30);let error;try{await p}catch(e){error=e.message}const next=await S.Engine.anonymizeText('Иванов Иван Иванович');return {error,text:next.text,stats:S.Background.stats}}''')
        self.assertIn('отменена',result['error']);self.assertIn('Сотрудник_',result['text']);self.assertGreater(result['stats']['aborted'],0)

    def test_legacy_05_word_key_remains_supported(self):
        old=self.browser.new_page()
        try:
            old.goto(self.url.replace('v0.6.0','v0.5.0'))
            data=old.evaluate('''async()=>{const S=SafeCycle,b=await S.Docx.load(await (await S.Docx.demo()).arrayBuffer()),out=await S.Docx.anonymize(b,'old.docx',{},false);return {key:out.key,data:await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(out.blob)})}}''')
        finally:old.close()
        restored=self.page.evaluate('''async ({key,data})=>{const out=await SafeCycle.Engine.restore(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,key);return {restored:out.restored,sameHash:out.sameHash}}''',data)
        self.assertTrue(restored['sameHash']);self.assertGreater(restored['restored'],0)

    def test_ui_csv_password_download_and_restore(self):
        path=self.work/'ui.csv';path.write_bytes(('ФИО;Код\n'+NAME+';00123\n').encode('cp1251'))
        self.page.locator('#file').set_input_files(path);self.page.locator('#config-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#process').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.assertTrue(self.page.locator('#key-encryption').is_checked());self.page.locator('#key-password').fill(PASSWORD);self.page.locator('#ack').check()
        with self.page.expect_download() as d:self.page.locator('#download-file').click()
        anon=self.work/'ui_anon.csv';d.value.save_as(anon);self.assertEqual(d.value.suggested_filename,'ui_anon.csv')
        with self.page.expect_download() as d:self.page.locator('#download-key').click()
        key=self.work/'ui_key.json';d.value.save_as(key);self.assertEqual(json.loads(key.read_text())['schema'],'bezopasnyj-cikl-encrypted-key');self.assertNotIn(NAME,key.read_text())
        self.page.locator('#restore-tab').click();self.page.locator('#key-file').set_input_files(key);self.page.locator('#restore-password').fill(PASSWORD);self.page.locator('#return-file').set_input_files(anon);self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#restore').click();self.page.locator('#download-restored').wait_for(state='visible')
        with self.page.expect_download() as d:self.page.locator('#download-restored').click()
        back=self.work/'ui-restored.csv';d.value.save_as(back);self.assertEqual(back.read_bytes(),path.read_bytes())
        self.page.locator('#anon-tab').click();self.page.locator('#reset').click();self.assertEqual(self.page.locator('#key-password').input_value(),'');self.assertEqual(self.page.locator('#restore-password').input_value(),'')

    def test_ui_batch_columns_shared_header_and_download_restore(self):
        paths=[self.work/'ui-a.xlsx',self.work/'ui-b.xlsx',self.work/'ui-word.docx'];workbook(paths[0]);workbook(paths[1]);stage5.fixture(paths[2])
        self.page.locator('#file').set_input_files(paths);self.page.locator('#batch-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden');self.assertEqual(self.page.locator('#sheets .sheet').count(),2)
        self.page.locator('#sheets select[data-mode]').first.select_option('scan');self.page.locator('#apply-matching').click();self.assertEqual(self.page.locator('#sheets select[data-mode]').nth(5).input_value(),'scan')
        self.page.locator('#process-batch').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden');self.page.locator('#ack').check();self.page.locator('#key-encryption').uncheck()
        with self.page.expect_download() as d:self.page.locator('#download-file').click()
        zipped=self.work/'ui-package.zip';d.value.save_as(zipped)
        with zipfile.ZipFile(zipped) as z:
            self.assertEqual(len(z.namelist()),3);z.extractall(self.work/'ui-package')
        with self.page.expect_download() as d:self.page.locator('#download-key').click()
        key=self.work/'ui-package-key.json';d.value.save_as(key);self.assertEqual(json.loads(key.read_text())['schema'],'bezopasnyj-cikl-batch-key')
        self.page.locator('#restore-tab').click();self.page.locator('#key-file').set_input_files(key);self.page.locator('#return-file').set_input_files(list((self.work/'ui-package').iterdir()));self.page.locator('#restore-options').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#restore').click();self.page.locator('#download-restored').wait_for(state='visible')
        with self.page.expect_download() as d:self.page.locator('#download-restored').click()
        zipback=self.work/'ui-back.zip';d.value.save_as(zipback)
        with zipfile.ZipFile(zipback) as z:self.assertEqual(len(z.namelist()),3);self.assertNotIn('ui-a_anon_restored.xlsx',z.namelist());self.assertIn('ui-a_restored.xlsx',z.namelist())

    def test_package_irreversible_employee_removed_from_word_and_key(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Csv.load(new TextEncoder().encode('ФИО\\nИванов Иван Иванович\\n').buffer),w=await S.Docx.load(await (await S.Docx.demo()).arrayBuffer()),items=[{id:'csv',name:'erase.csv',format:'csv',book:b},{id:'word',name:'memo.docx',format:'docx',book:w}],out=await S.Batch.process(items,{csv:[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'delete',category:'Сотрудник'}]}]},{dictionary:[{value:'Иванов Иван Иванович',category:'Сотрудник'}]},false),back=await S.Batch.restoreOne(await out.outputs[1].blob.arrayBuffer(),out.key,{fileName:'memo.docx',format:'docx'}),z=await S.Zip.Archive.read(await back.blob.arrayBuffer());return {key:out.key,back:await z.text('word/document.xml'),audit:out.audit}}''')
        self.assertNotIn(NAME,json.dumps(result['key'],ensure_ascii=False));self.assertNotIn('Иванову Ивану Ивановичу',json.dumps(result['key'],ensure_ascii=False));self.assertNotIn(NAME,result['back']);self.assertNotIn('Иванову Ивану Ивановичу',result['back']);self.assertFalse(result['audit']['blocked'],result['audit']['findings'])

    def test_conflicting_irreversible_and_numeric_modes_block_package(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Csv.load(new TextEncoder().encode('Сумма\\n175123\\n').buffer),c=await S.Csv.load(new TextEncoder().encode('Сумма\\n175123\\n').buffer),settings={a:[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'range',category:'Сумма'}]}],b:[{path:c.sheets[0].path,headerRow:1,columns:[{mode:'scale',category:'Сумма'}]}]};try{await S.Batch.process([{id:'a',name:'one.csv',format:'csv',book:b},{id:'b',name:'two.csv',format:'csv',book:c}],settings,{},false);return null}catch(e){return e.message}}''')
        self.assertIn('Согласуйте режимы',result)

    def test_csv_edited_columns_new_calculation_numeric_return(self):
        original=('ФИО;Сумма\n'+NAME+';150000,50\n').encode();out=self.csv_process(original,[{'mode':'alias','category':'Сотрудник'},{'mode':'scale','category':'Сумма'}])
        rows=list(csv.reader(io.StringIO(base64.b64decode(out['anon']).decode()),delimiter=';'));amount=float(rows[1][1]);edited=('Итог;Сумма;Имя\n'+str(amount*2)+';'+str(amount)+';'+rows[1][0]+'\n').encode()
        result=self.page.evaluate('''async ({data,key})=>{const p=key.transforms.columns[0],out=await SafeCycle.Engine.restore(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,key,null,null,{transforms:[{sheet:'CSV',column:0,sourceId:p.id},{sheet:'CSV',column:1,sourceId:p.id}]});return {text:await out.blob.text(),sameHash:out.sameHash,transformed:out.transformed}}''',{'data':base64.b64encode(edited).decode(),'key':out['key']})
        restored=list(csv.reader(io.StringIO(result['text']),delimiter=';'));self.assertAlmostEqual(float(restored[1][0]),300001.00,delta=.03);self.assertAlmostEqual(float(restored[1][1]),150000.50,delta=.02);self.assertEqual(restored[1][2],NAME);self.assertFalse(result['sameHash']);self.assertEqual(result['transformed'],2)

    def test_unknown_part_checked_against_values_from_other_file(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Csv.load(new TextEncoder().encode('Клиент\\nСеверный ветер\\n').buffer),w=await S.Docx.load(await (await S.Docx.demo()).arrayBuffer());w.archive.set('word/extra.xml','<unknown><value>Северный ветер</value></unknown>');w.docs.set('word/extra.xml',S.Xml.parse('<unknown><value>Северный ветер</value></unknown>'));const out=await S.Batch.process([{id:'csv',name:'one.csv',format:'csv',book:b},{id:'doc',name:'two.docx',format:'docx',book:w}],{csv:[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'alias',category:'Клиент'}]}]}, {},false);return out.audit}''')
        self.assertTrue(result['blocked']);self.assertTrue(any(f['file']=='two.docx' and f['part']=='word/extra.xml' for f in result['findings']))

    def test_macro_project_bytes_preserved_and_signed_project_rejected(self):
        source=self.work/'macro-base.xlsx';workbook(source);parts=stage1.parts(source)
        types=parts['[Content_Types].xml'].decode().replace('spreadsheetml.sheet.main+xml','ms-excel.sheet.macroEnabled.main+xml')
        rels=parts['xl/_rels/workbook.xml.rels'].decode().replace('</Relationships>','<Relationship Id="macro" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
        # Synthetic opaque bytes test preservation only; this is not a runnable VBA project.
        vba=b'SYNTHETIC-OPAQUE-VBA-CONTENT\0\x01';parts.update({'[Content_Types].xml':types.encode(),'xl/_rels/workbook.xml.rels':rels.encode(),'xl/vbaProject.bin':vba})
        macro=self.work/'macro.xlsm'
        with zipfile.ZipFile(macro,'w') as z:
            for name,data in parts.items():z.writestr(name,data)
        out=self.page.evaluate('''async data=>{const S=SafeCycle,b=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),cfg=b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Сотрудник'}))})),out=await S.Engine.anonymize(b,cfg,'macro.xlsm',false),z=await S.Zip.Archive.read(await out.blob.arrayBuffer());return {format:out.key.file.format,audit:out.audit,vba:[...await z.get('xl/vbaProject.bin')]}}''',base64.b64encode(macro.read_bytes()).decode())
        self.assertEqual(bytes(out['vba']),vba);self.assertEqual(out['format'],'xlsm');self.assertTrue(out['audit']['blocked']);self.assertTrue(any(f['part']=='xl/vbaProject.bin' for f in out['audit']['findings']))
        parts['xl/vbaProjectSignature.bin']=b'signed'
        with zipfile.ZipFile(macro,'w') as z:
            for name,data in parts.items():z.writestr(name,data)
        error=self.page.evaluate('''async data=>{try{await SafeCycle.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);return null}catch(e){return e.message}}''',base64.b64encode(macro.read_bytes()).decode());self.assertIn('подписаны',error)

    def test_xlsx_key_surrogate_boundary_keeps_complete_json(self):
        result=self.page.evaluate('''async()=>{const S=SafeCycle,key=(await S.Engine.anonymizeText('Иванов Иван Иванович')).key;key.extra='x'.repeat(29900)+'🧑'.repeat(30000);const blob=await S.ProtectedKey.save(key,null,'xlsx'),back=await S.ProtectedKey.open(await blob.arrayBuffer());return {equal:JSON.stringify(key)===JSON.stringify(back),length:back.extra.length}}''');self.assertTrue(result['equal']);self.assertEqual(result['length'],89900)

    def test_builtin_reliability_checks_and_all_checks_ui(self):
        checks=self.page.evaluate('async()=>await SafeCycle.ReliabilitySelfTest.run()');self.assertEqual(len(checks),4);self.assertTrue(all(c['pass'] for c in checks),checks)
        self.page.get_by_text('Для службы ИБ · ограничения и самопроверка',exact=True).click();self.page.locator('#selftest').click();self.page.locator('#busy').wait_for(state='hidden');self.assertEqual(self.page.locator('#tests .test-pass').count(),29);self.assertEqual(self.page.locator('#tests .test-fail').count(),0)

    def test_human_key_and_package_outputs_open_in_libreoffice(self):
        result,_=self.bundle(include_csv=False)
        data=self.page.evaluate('''async key=>{const blob=await SafeCycle.ProtectedKey.save(key,null,'xlsx');return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}''',result['key'])
        key=self.work/'lo-key.xlsx';key.write_bytes(base64.b64decode(data));dest=self.work/'stage6-lo';dest.mkdir(exist_ok=True);exe=shutil.which('soffice');self.assertIsNotNone(exe)
        for extension,files in [('xlsx',[self.work/'anon-staff.xlsx',self.work/'back-leaves.xlsx',key]),('docx',[self.work/'anon-memo.docx',self.work/'back-memo.docx'])]:
            run=subprocess.run([exe,f'-env:UserInstallation={self.work.as_uri()}/stage6-profile','--headless','--convert-to',extension,'--outdir',str(dest),*map(str,files)],capture_output=True,text=True,timeout=60);self.assertEqual(run.returncode,0,run.stdout+run.stderr)
            for f in files:self.assertTrue((dest/f.name).exists(),run.stdout+run.stderr)
        wb=openpyxl.load_workbook(dest/key.name);self.assertIn('Соответствия',wb.sheetnames)
        opened=self.page.evaluate('''async data=>await SafeCycle.ProtectedKey.open(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer)''',base64.b64encode((dest/key.name).read_bytes()).decode());self.assertEqual(opened,result['key'])

    def test_csv_100000_rows_worker_progress_and_page_responsiveness(self):
        # A complete file pass, not a loop that imitates the implementation.
        result=self.page.evaluate('''async()=>{const S=SafeCycle,n=100000,bytes=new TextEncoder().encode('ФИО;Код\\n'+Array.from({length:n},(_,i)=>'Иванов Иван Иванович;'+String(i).padStart(6,'0')).join('\\n')+'\\n');let beats=0,last=performance.now(),maxGap=0;const interval=setInterval(()=>{const now=performance.now();maxGap=Math.max(maxGap,now-last);last=now;beats++},20),started=performance.now();try{const b=await S.Csv.load(bytes.buffer),cfg=[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'alias',category:'Сотрудник'},{mode:'keep',category:'Скрыто'}]}],progress=[],out=await S.Csv.anonymize(b,cfg,'large.csv',false,null,p=>progress.push(p)),data=S.CsvCodec.read(await out.blob.arrayBuffer());return {rows:data.rows.length,first:data.rows[1],last:data.rows.at(-1),beats,maxGap,ms:performance.now()-started,progress,stats:S.Background.stats,audit:out.audit.total,replaced:out.replaced}}finally{clearInterval(interval)}}''')
        self.assertEqual(result['rows'],100001);self.assertEqual(result['first'][1],'000000');self.assertEqual(result['last'][1],'099999');self.assertEqual(result['replaced'],100000);self.assertEqual(result['audit'],0)
        self.assertGreater(result['beats'],30);self.assertLess(result['maxGap'],1500,result);self.assertIn(100,result['progress']);self.assertGreater(result['stats']['plans'],0);self.assertGreater(result['stats']['audits'],0)
        print('100000 rows:',json.dumps({k:result[k] for k in ['ms','maxGap','beats']},ensure_ascii=False),flush=True)

    def test_large_native_xml_preserves_namespaces_comments_cdata_and_escaped_attributes(self):
        result=self.page.evaluate('''async()=>{const X=SafeCycle.Xml,ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main',rows=Array.from({length:12000},(_,i)=>'<s:row r="'+(i+1)+'"><s:c r="A'+(i+1)+'" t="inlineStr"><s:is><s:t><![CDATA[</s:sheetData> & value '+i+']]></s:t></s:is></s:c><!-- </s:sheetData> --></s:row>').join(''),xml='<?xml version="1.0"?><?keep instruction?><s:worksheet xmlns:s="'+ns+'" xmlns:x="urn:custom"><s:sheetData>'+rows+'</s:sheetData><x:after label="a &gt; b &amp; c"/></s:worksheet>',doc=await X.parseAsync(xml),serialized=await X.serializeAsync(doc),again=X.parse(serialized);return {rows:again.getElementsByTagNameNS(ns,'row').length,last:again.getElementsByTagNameNS(ns,'t')[11999].textContent,attribute:again.getElementsByTagNameNS('urn:custom','after')[0].getAttribute('label'),comments:[...again.getElementsByTagNameNS(ns,'row')].filter(n=>n.lastChild.nodeType===8).length,pi:again.firstChild.nodeType}}''')
        self.assertEqual(result,{'rows':12000,'last':'</s:sheetData> & value 11999','attribute':'a > b & c','comments':12000,'pi':7})

    def test_large_native_xml_rejects_malformed_tail_and_entities(self):
        errors=self.page.evaluate('''async()=>{const X=SafeCycle.Xml,prefix='<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>',rows='<row><c><v>123</v></c></row>'.repeat(45000),errors=[];for(const xml of [prefix+rows+'</sheetData></broken>',prefix+rows+'<row><c></row></sheetData></worksheet>','<!DOCTYPE worksheet [<!ENTITY x "secret">]>'+prefix+rows+'</sheetData></worksheet>'])try{await X.parseAsync(xml);errors.push(null)}catch(e){errors.push(e.message)}return errors}''')
        self.assertTrue(all(errors),errors)

    def test_large_native_clone_preserves_carriage_returns_and_cdata(self):
        result=self.page.evaluate('''async()=>{const X=SafeCycle.Xml,ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main',doc=X.parse('<worksheet xmlns="'+ns+'"><sheetData>'+'<row><c><is><t>value</t></is></c></row>'.repeat(5000)+'</sheetData></worksheet>'),t=doc.getElementsByTagNameNS(ns,'t')[0];t.replaceChildren(doc.createTextNode('first\\r\\nline'),doc.createCDATASection('raw\\rdata'));const cloned=await X.cloneAsync(doc),c=cloned.getElementsByTagNameNS(ns,'t')[0];return {rows:cloned.getElementsByTagNameNS(ns,'row').length,text:c.firstChild.data,cdata:c.lastChild.data,cdataType:c.lastChild.nodeType,independent:c!==t}}''')
        self.assertEqual(result,{'rows':5000,'text':'first\r\nline','cdata':'raw\rdata','cdataType':4,'independent':True})
