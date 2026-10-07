"""Проверки детекторов и интерфейса; здесь используются учебные данные."""
import base64
import io
import json
import shutil
import subprocess
from pathlib import Path
import unittest

import openpyxl
import xlsxwriter
import test_stage1 as stage1


class Stage2(unittest.TestCase):
    # Общий стенд браузера и независимых читателей, без повторного наследования тестов.
    setUpClass = classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass = classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp = stage1.Stage1.setUp
    tearDown = stage1.Stage1.tearDown

    def detector(self, text, dictionary=None):
        return self.page.evaluate('''({text,rows})=>SafeCycle.Detectors.find(text,
          {dictionary:SafeCycle.Detectors.compileDictionary(rows||[])})''',
                                  {'text': text, 'rows': dictionary})

    def test_names_six_cases_and_name_only_without_swallowing_context(self):
        families = [
            ['Иванов Иван Иванович', 'Иванова Ивана Ивановича', 'Иванову Ивану Ивановичу',
             'Иванова Ивана Ивановича', 'Ивановым Иваном Ивановичем', 'Иванове Иване Ивановиче'],
            ['Петрова Анна Сергеевна', 'Петровой Анны Сергеевны', 'Петровой Анне Сергеевне',
             'Петрову Анну Сергеевну', 'Петровой Анной Сергеевной', 'Петровой Анне Сергеевне'],
        ]
        self.assertGreaterEqual(self.page.evaluate('SafeCycle.Detectors.names.length'), 300)
        for forms in families:
            identities = []
            for text in forms:
                matches = self.detector(text)
                self.assertEqual(len(matches), 1, text)
                self.assertEqual(matches[0]['category'], 'Сотрудник')
                self.assertEqual((matches[0]['start'], matches[0]['end']), (0, len(text)))
                identities.append(matches[0]['identity'])
            self.assertEqual(len(set(identities)), 1, forms)
        text = 'Письмо Марии Ивановне. Позвонить Ивану Ивановичу.'
        matches = self.detector(text)
        self.assertEqual(len(matches), 2)
        self.assertEqual([text[m['start']:m['end']] for m in matches], ['Марии Ивановне', 'Ивану Ивановичу'])

    def test_initials_require_confirmation_and_are_consistent_in_both_orders(self):
        source = 'Иванов Иван Иванович; Иванову И.И.; И. И. Иванов; Иван Иванович; Иванов.'
        result = self.page.evaluate('''async text=>{
          const S=SafeCycle,cat=S.Detectors.catalog([text]),proposal=cat.suggestions[0];
          const unmerged=await S.Engine.anonymizeText(text,{},false);
          const merged=await S.Engine.anonymizeText(text,{merges:[{shortId:proposal.shortId,fullId:proposal.candidates[0].id}]},false);
          return {unmerged:unmerged.text,merged:merged.text,entries:merged.key.entries,proposal};
        }''', source)
        self.assertEqual(result['unmerged'].count('[Сотрудник_0001]'), 3)
        self.assertEqual(result['unmerged'].count('[Сотрудник_0002]'), 2)
        self.assertEqual(result['merged'].count('[Сотрудник_0001]'), 5)
        self.assertEqual(len(result['entries']), 1)

    def test_ambiguous_initials_not_automatically_merged(self):
        source = 'Иванов Иван Иванович; Иванов Игорь Ильич; Иванов И.И.; Иванов.'
        result = self.page.evaluate('''async text=>{
          const S=SafeCycle,cat=S.Detectors.catalog([text]),out=await S.Engine.anonymizeText(text,{},false);
          return {suggestions:cat.suggestions,key:out.key,text:out.text,warnings:out.warnings};
        }''', source)
        self.assertEqual(len(result['suggestions'][0]['candidates']), 2)
        self.assertEqual(len(result['key']['entries']), 3)
        self.assertTrue(result['warnings'])
        self.assertTrue(result['text'].endswith('Иванов.'))

    def test_checksums_valid_invalid_and_random_numbers(self):
        fixtures = [
            ('7707083893', '7707083894', 'ИНН'),
            ('500100732259', '500100732258', 'ИНН'),
            ('112-233-445 95', '112-233-445 94', 'СНИЛС'),
            ('4111 1111 1111 1111', '4111 1111 1111 1112', 'Карта'),
            ('1027700132195', '1027700132196', 'ОГРН'),
        ]
        # ОГРНИП: независимое вычисление контрольной цифры средствами Python.
        prefix = '32699001234567'
        valid = prefix + str(int(prefix) % 13 % 10)
        fixtures.append((valid, valid[:-1] + str((int(valid[-1]) + 1) % 10), 'ОГРНИП'))
        for valid, invalid, category in fixtures:
            with self.subTest(category=category):
                self.assertIn(category, [m['category'] for m in self.detector(valid)])
                self.assertNotIn(category, [m['category'] for m in self.detector(invalid)])
        self.assertFalse(self.detector('1234567890; 123456789012; 000000000000; 9999999999'))

    def test_contextual_bank_passport_contract_organization_dates_money_address(self):
        text = ('Паспорт: 12 34 567890; КПП: 990001001; БИК: 044525999; '
                'р/с 40702810900000000001; к/с 30101810400000000001; договор № TEST-42; '
                'счёт № INV-17; заказ № ORDER-2; спецификация № SPEC-4; ООО «Учебный партнёр»; '
                'дата рождения: 29.02.2000; дата рождения: 31.02.2000; '
                'г. Москва, ул. Учебная, д. 7, кв. 12; цена 10000; скидка 15%; 150 000 руб.; '
                'demo@example.test; +7 (999) 123-45-67')
        found = self.detector(text)
        kinds = {m['kind'] for m in found}
        self.assertTrue({'Паспорт', 'КПП', 'БИК', 'Банковский счёт', 'Номер документа',
                         'Организация', 'Дата рождения', 'Адрес', 'Денежная сумма', 'Процент',
                         'Email', 'Телефон'} <= kinds, kinds)
        self.assertEqual(sum(m['kind'] == 'Дата рождения' for m in found), 1)
        self.assertFalse(self.detector('990001001; 044525999; 40702810900000000001'))

    def test_dictionary_forms_boundaries_and_regex_characters(self):
        rows = [{'value': 'Ромашка', 'category': 'Клиент'},
                {'value': 'Северный ветер', 'category': 'Проект'},
                {'value': 'Ёлка', 'category': 'Продукт'},
                {'value': 'C++', 'category': 'Продукт'},
                {'value': 'код#42', 'category': 'Секрет', 'exact': True}]
        text = 'Ромашки, РОМАШКОЙ, Северному ветру, ЕЛКОЙ, C++; код#42. Ромашкабанк.'
        found = self.detector(text, rows)
        self.assertEqual(len(found), 6, found)
        self.assertNotIn('Ромашкабанк', [text[m['start']:m['end']] for m in found])
        self.assertEqual(sum(m['category'] == 'Продукт' for m in found), 2)

    def test_overlapping_dictionary_fragment_does_not_expose_rest_of_person_or_email(self):
        source = 'Иванов Иван Иванович, demo@example.test.'
        rows = [{'value': 'Иван', 'category': 'Клиент'}, {'value': 'example', 'category': 'Проект'}]
        result = self.page.evaluate('''async ({text,rows})=>{
          const out=await SafeCycle.Engine.anonymizeText(text,{dictionary:rows},false);
          return out.text;
        }''', {'text': source, 'rows': rows})
        self.assertEqual(result, '[Сотрудник_0001], [Email_0001].')

    def test_scan_excel_cross_check_and_exact_roundtrip_rich_strings(self):
        source = self.work / 'scan.xlsx'
        wb = xlsxwriter.Workbook(source)
        ws = wb.add_worksheet('Данные')
        bold = wb.add_format({'bold': True})
        ws.write_row('A1', ['ФИО', 'Комментарий', 'Оставить', 'Клиент'])
        ws.write('A2', 'Иванов Иван Иванович')
        ws.write_rich_string('B2', 'Для ', bold, 'Иванову Ивану Ивановичу', '; demo@example.test; Ромашки.')
        ws.write('C2', 'Связаться: demo@example.test; повтор Иванов.')
        ws.write('D2', 'Ромашка')
        wb.close()
        payload = self.page.evaluate('''async data=>{
          const S=SafeCycle,buffer=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,book=await S.Xlsx.load(buffer);
          const settings=[{path:book.sheets[0].path,headerRow:1,columns:[
            {mode:'alias',category:'Сотрудник'},{mode:'scan',category:'Скрыто'},
            {mode:'keep',category:'Скрыто'},{mode:'alias',category:'Клиент'}]}];
          const out=await S.Engine.anonymize(book,settings,'scan.xlsx',false,undefined,undefined,
            {dictionary:[{category:'Клиент',value:'Ромашка'}]});
          const restored=await S.Engine.restore(await out.blob.arrayBuffer(),JSON.parse(JSON.stringify(out.key)));
          const a=await S.Xlsx.load(await out.blob.arrayBuffer());
          function b64(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}
          return {anon:await b64(out.blob),back:await b64(restored.blob),key:out.key,
            values:S.Xlsx.snapshot(a),cross:out.crossMatches};
        }''', base64.b64encode(source.read_bytes()).decode())
        anonymous = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['anon'])))
        self.assertEqual(anonymous['Данные']['A2'].value, 'Сотрудник_0001')
        self.assertIn('[Сотрудник_0001]', anonymous['Данные']['B2'].value)
        self.assertIn('[Сотрудник_0001]', anonymous['Данные']['C2'].value)
        self.assertIn('[Email_0001]', anonymous['Данные']['B2'].value)
        self.assertIn('[Email_0001]', anonymous['Данные']['C2'].value)
        self.assertIn('[Клиент_0001]', anonymous['Данные']['B2'].value)
        self.assertGreater(payload['cross'], 0)
        restored = openpyxl.load_workbook(io.BytesIO(base64.b64decode(payload['back'])), rich_text=True)
        original = openpyxl.load_workbook(source, rich_text=True)
        for row in original['Данные']:
            for cell in row:
                self.assertEqual(str(cell.value), str(restored['Данные'][cell.coordinate].value))
        self.assertEqual(str(original['Данные']['B2'].value), str(restored['Данные']['B2'].value))
        self.assertTrue(restored['Данные']['B2'].value[1].font.b)
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if executable:
            destination = self.work / 'stage2-lo'
            destination.mkdir(exist_ok=True)
            files = []
            for name, field in [('scan-anon.xlsx', 'anon'), ('scan-restored.xlsx', 'back')]:
                path = self.work / name
                path.write_bytes(base64.b64decode(payload[field]))
                files.append(path)
            run = subprocess.run([executable, f'-env:UserInstallation={self.work.as_uri()}/stage2-lo-profile',
                                  '--headless', '--convert-to', 'xlsx', '--outdir', str(destination),
                                  *map(str, files)], capture_output=True, text=True, timeout=60)
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            for path in files:
                converted = destination / path.name
                self.assertTrue(converted.exists(), run.stdout + run.stderr)
                openpyxl.load_workbook(converted)
            self.assertEqual(openpyxl.load_workbook(destination / files[1].name)['Данные']['B2'].value,
                             str(original['Данные']['B2'].value))


    def test_cross_check_can_be_disabled(self):
        result = self.page.evaluate('''async()=>{
          const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer());
          const settings=book.sheets.map(s=>({path:s.path,headerRow:1,
            columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Сотрудник'}))}));
          const out=await S.Engine.anonymize(book,settings,'demo.xlsx',false,undefined,undefined,{crosscheck:false});
          const anon=await S.Xlsx.load(await out.blob.arrayBuffer());
          return {text:S.Xlsx.value(anon.sheets[0].byRow.get(2).get(4),anon.strings),cross:out.crossMatches};
        }''')
        self.assertIn('Петровой Анной Сергеевной', result['text'])
        self.assertEqual(result['cross'], 0)

    def test_cross_check_clears_text_formula_cache_without_changing_source(self):
        source = self.work / 'text-formula.xlsx'
        wb = xlsxwriter.Workbook(source)
        ws = wb.add_worksheet('Данные')
        ws.write_row('A1', ['ФИО', 'Формула'])
        ws.write('A2', 'Иванов Иван Иванович')
        ws.write_formula('B2', '=A2', None, 'Иванов Иван Иванович')
        wb.close()
        result = self.page.evaluate('''async data=>{
          const S=SafeCycle,book=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);
          const settings=[{path:book.sheets[0].path,headerRow:1,columns:[
            {mode:'alias',category:'Сотрудник'},{mode:'keep',category:'Скрыто'}]}];
          let error;try{await S.Engine.anonymize(book,settings,'formula.xlsx',false)}catch(e){error=e.message}
          const out=await S.Engine.anonymize(book,settings,'formula.xlsx',false,undefined,undefined,{crosscheck:false});
          const anon=await S.Xlsx.load(await out.blob.arrayBuffer()),cell=anon.sheets[0].byRow.get(2).get(1);
          return {error,source:S.Xlsx.value(book.sheets[0].byRow.get(2).get(0),book.strings),
            value:S.Xlsx.value(cell,anon.strings),formula:S.Xml.children(cell,'f')[0]?.textContent};
        }''', base64.b64encode(source.read_bytes()).decode())
        self.assertIsNone(result.get('error'))
        self.assertEqual(result['source'], 'Иванов Иван Иванович')
        self.assertEqual(result['value'], '')
        self.assertEqual(result['formula'], 'A2')

    def test_legacy_stage1_key_is_supported(self):
        old = self.browser.new_page()
        try:
            old.goto(self.url.replace('v0.6.0', 'v0.1.0'))
            payload = old.evaluate('''async()=>{
              const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer());
              const settings=book.sheets.map(s=>({path:s.path,headerRow:1,
                columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Сотрудник'}))}));
              const out=await S.Engine.anonymize(book,settings,'legacy.xlsx',false);
              const base64=await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(out.blob)});
              return {base64,key:out.key};
            }''')
        finally:
            old.close()
        result = self.page.evaluate('''async({base64,key})=>{
          const S=SafeCycle,out=await S.Engine.restore(Uint8Array.from(atob(base64),c=>c.charCodeAt(0)).buffer,key);
          const book=await S.Xlsx.load(await out.blob.arrayBuffer());
          return {sameHash:out.sameHash,first:S.Xlsx.value(book.sheets[0].byRow.get(2).get(0),book.strings)};
        }''', payload)
        self.assertEqual(payload['key']['version'], '0.1.0')
        self.assertTrue(result['sameHash'])
        self.assertEqual(result['first'], 'Иванов Иван Иванович')

    def test_known_identifiers_keep_one_label_across_number_formats(self):
        result = self.page.evaluate('''()=>{
          const S=SafeCycle,rows=[
            ['8 (999) 123-45-67','alias','Телефон'],['Звонить +7 999 123 45 67.','keep','Скрыто'],
            ['79991234567','alias','Телефон'],['+7 (999) 123-45-67','scan','Скрыто'],
            ['112-233-445 95','alias','СНИЛС'],['СНИЛС 11223344595.','keep','Скрыто'],
            ['4111 1111 1111 1111','alias','Карта'],['Карта 4111111111111111.','keep','Скрыто']];
          const plan=S.Analyzer.plan(rows.map(([value,mode,category],i)=>({id:String(i),sheet:'Данные',ref:'A'+(i+2),
            column:0,value,mode,category,isText:true,type:'inlineStr'})),{},false);
          return {entries:plan.entries.map(e=>e.label),values:rows.map((_,i)=>plan.jobs.get(String(i))?.anonValue)};
        }''')
        self.assertCountEqual(result['entries'], ['Телефон_0001','СНИЛС_0001','Карта_0001'])
        self.assertEqual(result['values'][0], 'Телефон_0001')
        self.assertIn('[Телефон_0001]', result['values'][3])
        self.assertIn('[СНИЛС_0001]', result['values'][5])
        self.assertIn('[Карта_0001]', result['values'][7])

    def test_plain_text_exact_and_edited_restoration_unknown_label_report(self):
        source = ' Ёлкин Иван Иванович\nОтправить Ёлкину Ивану Ивановичу: demo@example.test. '
        result = self.page.evaluate('''async text=>{
          const S=SafeCycle,out=await S.Engine.anonymizeText(text,{},false),key=JSON.parse(JSON.stringify(out.key));
          const exact=await S.Engine.restoreText(out.text,key);
          const edited=await S.Engine.restoreText('Ответ: '+out.text+' [Сотрудник_9999]',key);
          return {out:out.text,exact,edited};
        }''', source)
        self.assertEqual(result['exact']['text'], source)
        self.assertTrue(result['exact']['sameHash'])
        self.assertFalse(result['edited']['sameHash'])
        self.assertEqual(result['edited']['unknown'], 1)
        self.assertIn('Ёлкин Иван Иванович', result['edited']['text'])
        self.assertNotIn('demo@example.test', result['out'])

    def test_dictionary_txt_and_xlsx_roundtrip_without_extra_hidden_sheets(self):
        rows = [{'value': 'Ромашка', 'category': 'Клиент', 'exact': False},
                {'value': 'код#42', 'category': 'Секрет', 'exact': True}]
        result = self.page.evaluate('''async rows=>{
          const D=SafeCycle.DictionaryFile,blob=await D.toXlsx(rows);
          const base64=await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)});
          return {txt:D.fromText(D.toText(rows)),xlsx:await D.fromXlsx(await blob.arrayBuffer()),base64,
            plain:D.fromText(['Ромашка','Вектор'].join(String.fromCharCode(10)),'Проект')};
        }''', rows)
        self.assertEqual(result['txt'], rows)
        self.assertEqual(result['xlsx'], rows)
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(result['base64'])))
        self.assertEqual(wb.sheetnames, ['Словарь'])
        self.assertEqual(wb['Словарь']['B2'].value, 'Ромашка')
        self.assertEqual(result['plain'][0]['category'], 'Проект')
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if executable:
            source = self.work / 'dictionary.xlsx'
            source.write_bytes(base64.b64decode(result['base64']))
            destination = self.work / 'dictionary-lo'
            destination.mkdir(exist_ok=True)
            run = subprocess.run([executable, f'-env:UserInstallation={self.work.as_uri()}/dictionary-lo-profile',
                                  '--headless', '--convert-to', 'xlsx', '--outdir', str(destination), str(source)],
                                 capture_output=True, text=True, timeout=60)
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            converted = destination / source.name
            self.assertTrue(converted.exists(), run.stdout + run.stderr)
            self.assertEqual(openpyxl.load_workbook(converted)['Словарь']['B2'].value, 'Ромашка')


    def test_ui_plaintext_key_download_and_restore(self):
        source = 'Иванов Иван Иванович: demo@example.test.'
        self.page.locator('#source-text').click()
        self.page.locator('#raw-text').fill(source)
        self.page.locator('#process-text').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertNotIn('Иванов', self.page.locator('#anon-text').input_value())
        self.assertTrue(self.page.locator('#copy-text').is_disabled())
        self.page.locator('#ack').check()
        anonymous = self.page.locator('#anon-text').input_value()
        self.assertTrue(self.page.locator('#copy-text').is_enabled())
        self.page.locator('#key-encryption').uncheck()
        with self.page.expect_download() as downloaded:
            self.page.locator('#download-key').click()
        key = self.work / 'text-key.json'
        downloaded.value.save_as(key)
        self.page.locator('#restore-tab').click()
        self.page.locator('#return-text-mode').click()
        self.page.locator('#key-file').set_input_files(key)
        self.page.locator('#return-text').fill(anonymous)
        self.page.locator('#restore').click()
        self.page.locator('#restored-text-field').wait_for(state='visible')
        self.assertEqual(self.page.locator('#restored-text').input_value(), source)
        self.assertIn('Отпечаток совпал', self.page.locator('#restore-report').inner_text())

    def test_ui_merge_requires_explicit_selection(self):
        self.page.locator('#source-text').click()
        self.page.locator('#raw-text').fill('Иванов Иван Иванович; Иванов И.И.')
        self.page.locator('#merge-refresh').click()
        selector = self.page.locator('#merge-list select')
        self.assertEqual(selector.input_value(), '')
        option = selector.locator('option').nth(1).get_attribute('value')
        selector.select_option(option)
        self.page.locator('#process-text').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertEqual(self.page.locator('#anon-text').input_value(), '[Сотрудник_0001]; [Сотрудник_0001]')

    def test_ui_dictionary_categories_and_manual_fragment_are_applied(self):
        self.page.locator('#source-text').click()
        self.page.locator('#dictionary-paste').fill('Ромашка\nОрион')
        self.page.locator('#dictionary-add').click()
        self.page.get_by_label('Категория словаря 2', exact=True).select_option('Проект')
        self.page.get_by_text('Скрыть фрагмент вручную', exact=True).click()
        self.page.locator('#manual-value').fill('код#42')
        self.page.locator('#manual-label').fill('Секрет')
        self.page.locator('#manual-add').click()
        self.page.locator('#raw-text').fill('Для Ромашки: Ориону, код#42.')
        self.page.locator('#process-text').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertEqual(self.page.locator('#anon-text').input_value(), 'Для [Клиент_0001]: [Проект_0001], [Секрет_0001].')

    def test_all_detector_selfchecks_pass(self):
        result = self.page.evaluate('async()=>await SafeCycle.DetectorSelfTest.run()')
        self.assertEqual(len(result), 6)
        self.assertTrue(all(test['pass'] for test in result), result)


if __name__ == '__main__':
    unittest.main(verbosity=2)
