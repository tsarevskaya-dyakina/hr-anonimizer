"""Word: независимые документы python-docx, ZIP/XML, браузер и LibreOffice. Данные вымышленные."""
import base64
import io
import json
from pathlib import Path
import shutil
import subprocess
import unittest
import zipfile
from lxml import etree as ET

from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import expect
import test_stage1 as stage1

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
NAME='Иванов Иван Иванович'
EMAIL='ivan@example.test'


def xml_text(payload):
    return ''.join(ET.fromstring(payload).itertext())


def rewrite(path, changes):
    data=stage1.parts(path);data.update(changes)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for p,b in data.items():z.writestr(p,b)


def fixture(path, images=False):
    doc=Document()
    doc.add_heading('Служебная записка',0)
    para=doc.add_paragraph(style='Normal')
    para.add_run('Ива').bold=True
    para.add_run('нов Иван ').italic=True
    para.add_run('Иванович согласовал документ.').underline=True
    doc.add_paragraph('Передать Иванову Ивану Ивановичу, e-mail '+EMAIL+'.')
    doc.add_paragraph('ООО «Учебный партнёр», ИНН 7707083893, КПП 773601001. Договор № TEST-42 на 150 000 руб.')
    doc.add_paragraph('Клиент Северный ветер. Отдел Звездоград, код 876543210987.')
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text=NAME;table.cell(0,1).text=EMAIL
    doc.sections[0].header.paragraphs[0].text='Контакт: '+NAME
    doc.sections[0].footer.paragraphs[0].text='Почта: '+EMAIL
    rev=doc.add_paragraph()
    deleted=OxmlElement('w:del');deleted.set(qn('w:id'),'1');deleted.set(qn('w:author'),'Удалённый автор');deleted.set(qn('w:date'),'2026-10-07T12:00:00Z')
    run=OxmlElement('w:r');t=OxmlElement('w:delText');t.text='Старый секрет old@example.test';run.append(t);deleted.append(run);rev._p.append(deleted)
    ins=OxmlElement('w:ins');ins.set(qn('w:id'),'2');ins.set(qn('w:author'),NAME);ins.set(qn('w:date'),'2026-10-07T12:00:00Z')
    run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Новая редакция: '+NAME;run.append(t);ins.append(run);rev._p.append(ins)
    commentp=doc.add_paragraph('Комментарий привязан здесь.')
    for tag in ['commentRangeStart','commentRangeEnd']:
        n=OxmlElement('w:'+tag);n.set(qn('w:id'),'0');commentp._p.append(n)
    cr=OxmlElement('w:r');n=OxmlElement('w:commentReference');n.set(qn('w:id'),'0');cr.append(n);commentp._p.append(cr)
    fieldp=doc.add_paragraph()
    for typ,value in [('begin',None),(None,' MERGEFIELD "'+EMAIL+'" '),('separate',None),(None,NAME),('end',None)]:
        r=OxmlElement('w:r')
        if typ:n=OxmlElement('w:fldChar');n.set(qn('w:fldCharType'),typ)
        else:n=OxmlElement('w:instrText' if 'MERGEFIELD' in value else 'w:t');n.text=value
        r.append(n);fieldp._p.append(r)
    simple=OxmlElement('w:fldSimple');simple.set(qn('w:instr'),'HYPERLINK "https://private.test/'+EMAIL+'"');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Контакт: '+NAME;r.append(t);simple.append(r);doc.add_paragraph()._p.append(simple)
    # VML text box: paragraph nested inside a containing paragraph.
    box=OxmlElement('w:txbxContent');bp=OxmlElement('w:p');br=OxmlElement('w:r');bt=OxmlElement('w:t');bt.text='Надпись: '+NAME;br.append(bt);bp.append(br);box.append(bp)
    shape=ET.Element('{urn:schemas-microsoft-com:vml}shape',nsmap={'v':'urn:schemas-microsoft-com:vml'})
    shape.set('id','TextBox1');shape.set('style','width:200pt;height:40pt')
    textbox=ET.SubElement(shape,'{urn:schemas-microsoft-com:vml}textbox');textbox.append(box)
    pic=OxmlElement('w:pict');pic.append(shape);r=OxmlElement('w:r');r.append(pic);doc.add_paragraph()._p.append(r)
    if images:
        png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aN1sAAAAASUVORK5CYII=')
        doc.add_picture(io.BytesIO(png),width=Inches(.2))
    doc.core_properties.author=NAME;doc.core_properties.last_modified_by=NAME;doc.core_properties.title='Личная записка '+NAME
    doc.save(path)
    data=stage1.parts(path)
    comments=f'<w:comments xmlns:w="{W}"><w:comment w:id="0" w:author="{NAME}" w:initials="ИИ"><w:p><w:r><w:t>От {NAME}: {EMAIL}</w:t></w:r></w:p></w:comment></w:comments>'
    notes={}
    for filename,root,child in [('footnotes','footnotes','footnote'),('endnotes','endnotes','endnote')]:
        notes['word/'+filename+'.xml']=f'<w:{root} xmlns:w="{W}"><w:{child} w:id="1"><w:p><w:r><w:t>Примечание: {NAME}, {EMAIL}</w:t></w:r></w:p></w:{child}></w:{root}>'
    rel=ET.fromstring(data['word/_rels/document.xml.rels'])
    for name in ['comments','footnotes','endnotes']:
        ET.SubElement(rel,'{'+P+'}Relationship',Id='extra-'+name,Type=R+'/'+name,Target=name+'.xml')
    types=ET.fromstring(data['[Content_Types].xml'])
    for name in ['comments','footnotes','endnotes']:
        ET.SubElement(types,'{'+CT+'}Override',PartName='/word/'+name+'.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.'+name+'+xml')
    # Actual references in the body for the notes.
    root=ET.fromstring(data['word/document.xml']);body=root.find('{'+W+'}body')
    n=ET.Element('{'+W+'}p')
    for name in ['footnoteReference','endnoteReference']:
        r=ET.SubElement(n,'{'+W+'}r');ET.SubElement(r,'{'+W+'}'+name,{'{'+W+'}id':'1'})
    body.insert(len(body)-1,n)
    rewrite(path,{**notes,'word/comments.xml':comments,'word/_rels/document.xml.rels':ET.tostring(rel),'[Content_Types].xml':ET.tostring(types),'word/document.xml':ET.tostring(root),'customXml/secret.xml':'<secret>old@example.test</secret>'})


class Stage5(unittest.TestCase):
    setUpClass=classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass=classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp=stage1.Stage1.setUp
    tearDown=stage1.Stage1.tearDown

    def process(self, source=None, options=None, restore=True):
        source=source or self.work/'word.docx'
        if not source.exists():fixture(source)
        result=self.page.evaluate('''async ({data,options,restore})=>{const S=SafeCycle,b=await S.Docx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),out=await S.Docx.anonymize(b,'word.docx',{dictionary:[{category:'Клиент',value:'Северный ветер'}],...(options||{})},false),back=restore?await S.Engine.restore(await out.blob.arrayBuffer(),JSON.parse(JSON.stringify(out.key))):null;async function b64(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}return {key:out.key,audit:out.audit,summary:b.summary,preview:out.preview,anon:await b64(out.blob),back:back?await b64(back.blob):null,restored:back?back.restored:null}}''',{'data':base64.b64encode(source.read_bytes()).decode(),'options':options or {},'restore':restore})
        anon=self.work/'word-anon.docx';anon.write_bytes(base64.b64decode(result['anon']))
        if result['back']:(self.work/'word-back.docx').write_bytes(base64.b64decode(result['back']))
        return result,anon

    def test_all_word_stories_fields_revisions_comments_and_properties(self):
        result,anon=self.process()
        self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        parts=stage1.parts(anon)
        self.assertNotIn('word/comments.xml',parts)
        self.assertNotIn('customXml/secret.xml',parts)
        for part,data in parts.items():
            if part.endswith(('.xml','.rels','.vml')):
                for secret in [NAME,EMAIL,'Иванову Ивану Ивановичу','old@example.test','Удалённый автор','Северный ветер']:
                    self.assertNotIn(secret,data.decode(),part)
        main=ET.fromstring(parts['word/document.xml'])
        for name in ['ins','del','delText','fldChar','fldSimple','instrText','hyperlink','commentReference','commentRangeStart','commentRangeEnd']:
            self.assertEqual(main.findall('.//{'+W+'}'+name),[],name)
        d=Document(anon)
        self.assertEqual(len(d.tables),1)
        label=next(e['label'] for e in result['key']['entries'] if e['canonical']==NAME)
        for part in ['word/header1.xml','word/footnotes.xml','word/endnotes.xml','word/document.xml']:
            self.assertIn('['+label+']',xml_text(parts[part]),part)
        self.assertEqual(d.core_properties.author,'')
        self.assertTrue(d.paragraphs[1].runs[0].bold)
        self.assertTrue(d.paragraphs[1].runs[1].italic)
        self.assertEqual(d.paragraphs[1].runs[0].text,'['+label+']')
        self.assertIn(' согласовал документ.',d.paragraphs[1].text)

    def test_exact_restore_keeps_rich_runs_cases_and_does_not_resurrect_deleted_data(self):
        result,_=self.process()
        d=Document(self.work/'word-back.docx')
        self.assertEqual(d.paragraphs[1].runs[0].text,'Ива')
        self.assertTrue(d.paragraphs[1].runs[0].bold)
        self.assertEqual(d.paragraphs[1].runs[1].text,'нов Иван ')
        self.assertTrue(d.paragraphs[1].runs[1].italic)
        self.assertIn('Иванову Ивану Ивановичу',d.paragraphs[2].text)
        self.assertEqual(d.tables[0].cell(0,0).text,NAME)
        self.assertEqual(d.sections[0].header.paragraphs[0].text,'Контакт: '+NAME)
        back=stage1.parts(self.work/'word-back.docx')
        self.assertNotIn('word/comments.xml',back)
        self.assertNotIn('old@example.test',back['word/document.xml'].decode())
        self.assertEqual(d.core_properties.author,'')
        self.assertNotIn('old@example.test',json.dumps(result['key'],ensure_ascii=False))
        self.assertNotIn('Удалённый автор',json.dumps(result['key'],ensure_ascii=False))
        self.assertGreater(result['restored'],0)

    def test_keep_revisions_and_comments_masks_deleted_text_authors_and_initials(self):
        result,anon=self.process(options={'acceptChanges':False,'deleteComments':False})
        self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        data=stage1.parts(anon)
        main=ET.fromstring(data['word/document.xml'])
        self.assertTrue(main.findall('.//{'+W+'}del'))
        self.assertTrue(main.findall('.//{'+W+'}ins'))
        self.assertNotIn('old@example.test',data['word/document.xml'].decode())
        for n in main.findall('.//{'+W+'}del')+main.findall('.//{'+W+'}ins'):
            self.assertTrue(n.attrib['{'+W+'}author'].startswith('Сотрудник_'))
        comment=ET.fromstring(data['word/comments.xml'])[0]
        self.assertEqual(comment.attrib['{'+W+'}author'],comment.attrib['{'+W+'}initials'])
        back=stage1.parts(self.work/'word-back.docx')
        self.assertIn('old@example.test',back['word/document.xml'].decode())
        original=ET.fromstring(back['word/comments.xml'])[0]
        self.assertEqual(original.attrib['{'+W+'}author'],NAME)
        self.assertEqual(original.attrib['{'+W+'}initials'],'ИИ')

    def test_images_default_warn_and_block_optional_removal(self):
        source=self.work/'images.docx';fixture(source,images=True)
        result,anon=self.process(source)
        self.assertEqual(result['summary']['images'],1)
        self.assertTrue(result['audit']['blocked'])
        self.assertTrue(any(f['part'].startswith('word/media/') for f in result['audit']['findings']))
        self.assertEqual(len(Document(anon).inline_shapes),1)
        clean,anon=self.process(source,{'removeImages':True})
        self.assertFalse(clean['audit']['blocked'],clean['audit']['findings'])
        self.assertEqual(len(Document(anon).inline_shapes),0)
        self.assertFalse(any(p.startswith('word/media/') for p in stage1.parts(anon)))
        self.assertEqual(len(Document(self.work/'word-back.docx').inline_shapes),0)

    def test_unknown_xml_known_source_is_blocked_with_paragraph_location(self):
        source=self.work/'unknown-word.docx';fixture(source)
        rewrite(source,{'word/extension.xml':f'<w:extension xmlns:w="{W}"><w:p><w:r><w:t>{NAME}</w:t></w:r></w:p></w:extension>'})
        # Supported Word paragraphs in unusual parts are still processed.
        result,_=self.process(source)
        self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        rewrite(source,{'word/extension.xml':'<extra>'+NAME+'</extra>'})
        result,_=self.process(source)
        self.assertTrue(result['audit']['blocked'])
        self.assertTrue(any(f['part']=='word/extension.xml' for f in result['audit']['findings']))
        # Audit independently diagnoses fragments split over multiple Word runs.
        found=self.page.evaluate('''async key=>{const S=SafeCycle,a=new S.Zip.Archive();a.set('word/document.xml','<w:document xmlns:w="'+S.Docx.W+'"><w:body><w:p><w:r><w:t>Ива</w:t></w:r><w:r><w:t>нов Иван Иванович</w:t></w:r></w:p></w:body></w:document>');return S.Audit.scan(a,key)}''',result['key'])
        self.assertTrue(any(f.get('paragraph')==1 and f['match']==NAME for f in found['findings']),found)

    def test_edited_restore_finds_split_markers_new_paragraphs_and_unknown_labels(self):
        result,anon=self.process()
        d=Document(anon);label=next(e['label'] for e in result['key']['entries'] if e['canonical']==NAME)
        p=d.add_paragraph('Ответ ИИ: ');p.add_run('['+label[:8]).bold=True;p.add_run(label[8:]+']')
        d.add_paragraph('[Скрыто_9999]')
        edited=self.work/'edited-word.docx';d.save(edited)
        out=self.page.evaluate('''async ({data,key})=>{const S=SafeCycle,o=await S.Engine.restore(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,key);return {sameHash:o.sameHash,unknown:o.unknown,data:await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(o.blob)})}}''',{'data':base64.b64encode(edited.read_bytes()).decode(),'key':result['key']})
        self.assertFalse(out['sameHash']);self.assertEqual(out['unknown'],1)
        back=Document(io.BytesIO(base64.b64decode(out['data'])))
        self.assertEqual(back.paragraphs[-2].text,'Ответ ИИ: '+NAME)
        self.assertTrue(back.paragraphs[-2].runs[1].bold)
        self.assertEqual(back.paragraphs[-1].text,'[Скрыто_9999]')

    def test_dictionary_across_runs_preserves_unmatched_text_and_first_run_style(self):
        source=self.work/'dictionary-split.docx'
        d=Document();p=d.add_paragraph('Клиент: ');p.add_run('Север').bold=True;p.add_run('ный ветер').italic=True;p.add_run(' согласовал.');d.save(source)
        result,anon=self.process(source)
        self.assertFalse(result['audit']['blocked'])
        p=Document(anon).paragraphs[0]
        self.assertEqual(p.text,'Клиент: [Клиент_0001] согласовал.')
        self.assertEqual(p.runs[0].text,'Клиент: ');self.assertTrue(p.runs[1].bold);self.assertEqual(p.runs[3].text,' согласовал.')
        self.assertEqual(Document(self.work/'word-back.docx').paragraphs[0].runs[2].text,'ный ветер')

    def test_external_relationships_and_custom_properties_removed(self):
        source=self.work/'links.docx';fixture(source)
        data=stage1.parts(source)
        rel=ET.fromstring(data['word/_rels/document.xml.rels']);ET.SubElement(rel,'{'+P+'}Relationship',Id='private-link',Type=R+'/hyperlink',Target='https://private.test/'+EMAIL,TargetMode='External')
        root=ET.fromstring(data['word/document.xml']);p=root.find('.//{'+W+'}body/{'+W+'}p');link=ET.SubElement(p,'{'+W+'}hyperlink',{'{'+R+'}id':'private-link'});r=ET.SubElement(link,'{'+W+'}r');ET.SubElement(r,'{'+W+'}t').text=NAME
        rewrite(source,{'word/_rels/document.xml.rels':ET.tostring(rel),'word/document.xml':ET.tostring(root),'docProps/custom.xml':'<properties><secret>'+EMAIL+'</secret></properties>'})
        result,anon=self.process(source)
        self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        data=stage1.parts(anon)
        self.assertNotIn('docProps/custom.xml',data)
        self.assertNotIn('private.test',data['word/_rels/document.xml.rels'].decode())
        self.assertNotIn('hyperlink',data['word/document.xml'].decode())
        Document(anon)

    def test_cancel_keeps_source_and_does_not_publish_result(self):
        source=self.work/'cancel-word.docx';fixture(source)
        result=self.page.evaluate('''async data=>{const S=SafeCycle,b=await S.Docx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),before=await b.archive.text('word/document.xml'),ctrl=new AbortController();const p=S.Docx.anonymize(b,'cancel.docx',{},false,ctrl.signal);ctrl.abort();let error;try{await p}catch(e){error=e.message}return {error,unchanged:before===await b.archive.text('word/document.xml')}}''',base64.b64encode(source.read_bytes()).decode())
        self.assertIn('отменена',result['error']);self.assertTrue(result['unchanged'])

    def test_corrupt_word_key_is_rejected(self):
        result,_=self.process()
        for field,value in [('path',[-1]),('ns','bad'),('originalXml','<bad/>')]:
            key=json.loads(json.dumps(result['key']));key['docxRecords'][0][field]=value
            error=self.page.evaluate('''async ({data,key})=>{try{await SafeCycle.Engine.restore(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,key);return null}catch(e){return e.message}}''',{'data':result['anon'],'key':key})
            self.assertIsNotNone(error,field)

    def test_docx_key_restores_plain_text_ai_response(self):
        result,_=self.process();label=next(e['label'] for e in result['key']['entries'] if e['canonical']==NAME)
        back=self.page.evaluate('async ({key,text})=>await SafeCycle.Engine.restoreText(text,key)',{'key':result['key'],'text':'Ответ: ['+label+']'})
        self.assertEqual(back['text'],'Ответ: '+NAME)

    def test_demo_and_outputs_open_in_libreoffice(self):
        result,anon=self.process();files=[anon,self.work/'word-back.docx']
        data=self.page.evaluate('async()=>{const b=await SafeCycle.Docx.demo();return await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(",")[1]);r.readAsDataURL(b)})}')
        demo=self.work/'word-demo.docx';demo.write_bytes(base64.b64decode(data));Document(demo);files.append(demo)
        destination=self.work/'word-lo';destination.mkdir(exist_ok=True)
        executable=shutil.which('soffice');self.assertIsNotNone(executable)
        run=subprocess.run([executable,f'-env:UserInstallation={self.work.as_uri()}/word-profile','--headless','--convert-to','docx','--outdir',str(destination),*map(str,files)],capture_output=True,text=True,timeout=60)
        self.assertEqual(run.returncode,0,run.stdout+run.stderr)
        for source in files:
            converted=destination/source.name;self.assertTrue(converted.exists(),run.stdout+run.stderr);Document(converted)
        body='\n'.join(p.text for p in Document(destination/anon.name).paragraphs)
        self.assertIn('Сотрудник_',body);self.assertNotIn(NAME,body);self.assertNotIn('old@example.test',body)

    def test_ui_word_options_download_restore_and_reset(self):
        source=self.work/'ui-word.docx';fixture(source)
        self.page.locator('#source-word').click()
        self.page.locator('#file').set_input_files(source)
        self.page.locator('#word-options').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.assertTrue(self.page.locator('#accept-changes').is_checked());self.assertTrue(self.page.locator('#delete-comments').is_checked());self.assertFalse(self.page.locator('#remove-word-images').is_checked())
        self.assertIn('ui-word.docx',self.page.locator('#word-summary').inner_text())
        self.assertTrue(self.page.locator('#config-card').is_hidden())
        self.page.locator('#process-word').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#ack').check();expect(self.page.locator('#download-file')).to_be_enabled()
        with self.page.expect_download() as download:self.page.locator('#download-file').click()
        anon=self.work/'ui-anon.docx';download.value.save_as(anon);self.assertEqual(download.value.suggested_filename,'ui-word_anon.docx');Document(anon)
        self.page.locator('#key-encryption').uncheck()
        with self.page.expect_download() as download:self.page.locator('#download-key').click()
        key=self.work/'ui-word-key.json';download.value.save_as(key)
        self.page.locator('#restore-tab').click();self.page.locator('#key-file').set_input_files(key);self.page.locator('#return-file').set_input_files(anon)
        self.page.locator('#busy').wait_for(state='hidden');self.page.locator('#restore').click();self.page.locator('#download-restored').wait_for(state='visible')
        with self.page.expect_download() as download:self.page.locator('#download-restored').click()
        self.assertTrue(download.value.suggested_filename.endswith('_restored.docx'))
        restored=self.work/'ui-back.docx';download.value.save_as(restored);self.assertIn(NAME,Document(restored).paragraphs[1].text)
        self.page.locator('#anon-tab').click();self.page.locator('#reset').click();self.assertEqual(self.page.locator('#word-summary').text_content(),'');self.assertTrue(self.page.locator('#result-card').is_hidden())

    def test_ui_word_manual_hide_and_image_gate(self):
        source=self.work/'ui-image.docx';fixture(source,images=True)
        self.page.locator('#file').set_input_files(source);self.page.locator('#word-options').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#process-word').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#ack').check();expect(self.page.locator('#download-file')).to_be_disabled()
        self.page.locator('#remove-word-images').check();self.page.locator('#process-word').click();self.page.locator('#result-card').wait_for(state='visible');self.page.locator('#busy').wait_for(state='hidden')
        self.page.get_by_role('button',name='Скрыть Звездоград',exact=True).first.click();self.page.locator('#busy').wait_for(state='hidden')
        self.assertNotIn('Звездоград',self.page.locator('#manual-review').inner_text())
        self.page.locator('#ack').check();expect(self.page.locator('#download-file')).to_be_enabled()

    def test_people_namespaced_identifiers_removed_and_author_roundtrip(self):
        source=self.work/'people.docx';fixture(source)
        ns='http://schemas.microsoft.com/office/word/2012/wordml'
        people=f'<w15:people xmlns:w15="{ns}"><w15:person w15:author="{NAME}"><w15:presenceInfo w15:providerId="private-provider" w15:userId="hidden-identity@example.test"/></w15:person></w15:people>'
        data=stage1.parts(source);rels=ET.fromstring(data['word/_rels/document.xml.rels']);ET.SubElement(rels,'{'+P+'}Relationship',Id='people',Type='http://schemas.microsoft.com/office/2011/relationships/people',Target='people.xml')
        types=ET.fromstring(data['[Content_Types].xml']);ET.SubElement(types,'{'+CT+'}Override',PartName='/word/people.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml')
        rewrite(source,{'word/people.xml':people,'word/_rels/document.xml.rels':ET.tostring(rels),'[Content_Types].xml':ET.tostring(types)})
        default,anon=self.process(source);self.assertNotIn('word/people.xml',stage1.parts(anon));self.assertNotIn('hidden-identity@example.test',json.dumps(default['key']))
        kept,anon=self.process(source,{'acceptChanges':False,'deleteComments':False});parts=stage1.parts(anon)
        label=next(e['label'] for e in kept['key']['entries'] if e['canonical']==NAME)
        self.assertFalse(kept['audit']['blocked'],kept['audit']['findings'])
        node=ET.fromstring(parts['word/people.xml']);self.assertEqual(node.find('{'+ns+'}person').get('{'+ns+'}author'),label)
        self.assertNotIn('hidden-identity',parts['word/people.xml'].decode());self.assertNotIn('private-provider',parts['word/people.xml'].decode())
        self.assertEqual(ET.fromstring(stage1.parts(self.work/'word-back.docx')['word/people.xml']).find('{'+ns+'}person').get('{'+ns+'}author'),NAME)
        self.assertNotIn('hidden-identity',json.dumps(kept['key']))

    def test_nested_changed_paragraphs_restore_without_location_collision(self):
        source=self.work/'nested.docx';fixture(source)
        root=ET.fromstring(stage1.parts(source)['word/document.xml']);box=root.find('.//{'+W+'}txbxContent');outer=box
        while outer.tag!='{'+W+'}p':outer=outer.getparent()
        r=ET.Element('{'+W+'}r');ET.SubElement(r,'{'+W+'}t').text='Внешний контакт '+EMAIL;outer.insert(0,r)
        rewrite(source,{'word/document.xml':ET.tostring(root)})
        result,anon=self.process(source)
        self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        back=ET.fromstring(stage1.parts(self.work/'word-back.docx')['word/document.xml'])
        self.assertEqual(xml_text(ET.tostring(back.find('.//{'+W+'}txbxContent'))),'Надпись: '+NAME)
        self.assertIn('Внешний контакт '+EMAIL,xml_text(ET.tostring(back)))
        self.assertNotIn(NAME,xml_text(stage1.parts(anon)['word/document.xml']))

    def test_accept_deleted_rows_paragraph_breaks_and_old_property_history(self):
        source=self.work/'accepted.docx';fixture(source)
        root=ET.fromstring(stage1.parts(source)['word/document.xml']);body=root.find('{'+W+'}body')
        first=ET.Element('{'+W+'}p');prop=ET.SubElement(first,'{'+W+'}pPr');rp=ET.SubElement(prop,'{'+W+'}rPr');ET.SubElement(rp,'{'+W+'}del',{'{'+W+'}id':'42','{'+W+'}author':NAME})
        r=ET.SubElement(first,'{'+W+'}r');ET.SubElement(r,'{'+W+'}t').text='Иванов '
        second=ET.Element('{'+W+'}p');r=ET.SubElement(second,'{'+W+'}r');ET.SubElement(r,'{'+W+'}t').text='Иван Иванович'
        body.insert(1,first);body.insert(2,second)
        table=root.find('.//{'+W+'}tbl');row=ET.SubElement(table,'{'+W+'}tr');props=ET.SubElement(row,'{'+W+'}trPr');ET.SubElement(props,'{'+W+'}del',{'{'+W+'}id':'43','{'+W+'}author':'Старый автор'})
        cell=ET.SubElement(row,'{'+W+'}tc');p=ET.SubElement(cell,'{'+W+'}p');r=ET.SubElement(p,'{'+W+'}r');ET.SubElement(r,'{'+W+'}t').text='deleted-row@example.test'
        change=ET.SubElement(prop,'{'+W+'}pPrChange',{'{'+W+'}id':'44','{'+W+'}author':'Старый автор'});ET.SubElement(change,'{'+W+'}pPr')
        rewrite(source,{'word/document.xml':ET.tostring(root)})
        result,anon=self.process(source);self.assertFalse(result['audit']['blocked'],result['audit']['findings'])
        self.assertNotIn('deleted-row@example.test',json.dumps(result['key'],ensure_ascii=False));self.assertNotIn('Старый автор',json.dumps(result['key'],ensure_ascii=False))
        self.assertEqual(len(Document(anon).tables[0].rows),1)
        back=Document(self.work/'word-back.docx');self.assertEqual(back.paragraphs[1].text,NAME)
        self.assertNotIn('deleted-row',xml_text(stage1.parts(self.work/'word-back.docx')['word/document.xml']))

    def test_builtin_word_checks_and_ui_all_stage_checks(self):
        results=self.page.evaluate('async()=>await SafeCycle.WordSelfTest.run()')
        self.assertEqual(len(results),4);self.assertTrue(all(r['pass'] for r in results),results)
        self.page.get_by_text('Для службы ИБ · ограничения и самопроверка',exact=True).click()
        self.page.locator('#selftest').click();self.page.locator('#busy').wait_for(state='hidden')
        self.assertEqual(self.page.locator('#tests .test-pass').count(),29)
        self.assertEqual(self.page.locator('#tests .test-fail').count(),0)
        self.page.locator('#source-word').click();self.page.locator('#privacy-demo').click();self.page.locator('#busy').wait_for(state='hidden')
        self.assertTrue(self.page.locator('#config-card').is_visible());self.assertTrue(self.page.locator('#word-options').is_hidden())
