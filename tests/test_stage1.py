"""Независимые проверки первого этапа; все данные в примерах вымышленные."""
import base64
import functools
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import unittest
import zipfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import xlsxwriter
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / 'bezopasnyj-cikl-v1.0.0.html'


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def make_fixture(path):
    wb = xlsxwriter.Workbook(path)
    sheet = wb.add_worksheet('Сотрудники')
    bold = wb.add_format({'bold': True, 'bg_color': '#eaf4ed'})
    date = wb.add_format({'num_format': 'dd.mm.yyyy'})
    sheet.merge_range('A1:F1', 'Тестовая книга', bold)
    sheet.write_row('A3', ['ФИО', 'Табельный номер', 'Оклад', 'Дата', 'Секрет', 'Комментарий'], bold)
    rows = [
        [' Ёлкин Иван Иванович ', '00123', 80000, 45000, 'Стереть меня', '<img src=x onerror=alert(1)>'],
        ['елкин иван иванович', '00007', 90000.50, 45001, 'Безвозвратно', 'Комментарий оставлен'],
        ['Петрова Анна Сергеевна', '00000', 72000, 45002, 'Удалить значение', ''],
    ]
    for row_num, values in enumerate(rows, 3):
        sheet.write_row(row_num, 0, values)
        sheet.write_number(row_num, 3, values[3], date)
    sheet.write_formula('C7', '=SUM(C4:C6)', None, 242000.5)
    sheet.write_comment('A4', 'Этап 4: комментарий остаётся исходным', {'author': 'Автор теста'})
    sheet.set_row(4, None, None, {'hidden': True})
    sheet.set_column('A:A', 36)
    sheet.set_column('B:B', 18, None, {'hidden': True})
    sheet.freeze_panes(3, 1)
    sheet.autofilter('A3:F6')
    sheet.conditional_format('C4:C6', {'type': 'cell', 'criteria': '>', 'value': 85000, 'format': bold})
    sheet.data_validation('F4:F6', {'validate': 'list', 'source': ['Комментарий оставлен', 'Новый']})
    chart = wb.add_chart({'type': 'column'})
    chart.add_series({'values': '=Сотрудники!$C$4:$C$6', 'categories': '=Сотрудники!$A$4:$A$6'})
    sheet.insert_chart('H3', chart)
    hidden = wb.add_worksheet('Скрыто')
    hidden.very_hidden()
    hidden.write_row('A1', ['ФИО', 'Табельный номер'], bold)
    hidden.write_row('A2', [' Ёлкин Иван Иванович ', '00123'])
    wb.add_worksheet('Пустой лист')
    wb.set_properties({'author': 'Вымышленный автор'})
    wb.close()


def parts(path):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        return {name: archive.read(name) for name in archive.namelist()}


class Stage1(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix='safe-cycle-tests-')
        cls.work = Path(cls.temp.name)
        cls.source = cls.work / 'source.xlsx'
        make_fixture(cls.source)
        handler = functools.partial(QuietHandler, directory=str(ROOT))
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f'http://127.0.0.1:{cls.server.server_port}/{HTML.name}'
        cls.pw = sync_playwright().start()
        browser_path = os.environ.get('CHROMIUM_PATH') or shutil.which('chromium')
        options = {'headless': True}
        if browser_path:
            options['executable_path'] = browser_path
        cls.browser = cls.pw.chromium.launch(**options)

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.pw.stop()
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def setUp(self):
        self.page = self.browser.new_page(viewport={'width': 1440, 'height': 1000})
        self.errors = []
        self.console_errors = []
        self.requests = []
        self.worker_urls = []
        self.page.on('worker', lambda worker: self.worker_urls.append(worker.url))
        self.page.on('pageerror', lambda error: self.errors.append(str(error)))
        self.page.on('console', lambda message: self.console_errors.append(message.text) if message.type == 'error' else None)
        self.page.on('request', lambda request: self.requests.append(request.url))
        self.page.goto(self.url)

    def tearDown(self):
        self.assertEqual(self.errors, [], 'Ошибка JavaScript')
        self.assertEqual(self.console_errors, [], 'Ошибка в консоли браузера')
        self.assertEqual([url for url in self.requests if not url.startswith('blob:')], [self.url], 'Неожиданный сетевой запрос после открытия HTML')
        self.assertTrue(all(url in self.worker_urls for url in self.requests if url.startswith('blob:')), 'Неожиданный локальный ресурс')
        self.page.close()

    def output(self, mode='alias'):
        source64 = base64.b64encode(self.source.read_bytes()).decode()
        value = self.page.evaluate('''async ({data,mode})=>{
          const S=SafeCycle;
          const buffer=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer;
          const book=await S.Xlsx.load(buffer);
          const settings=book.sheets.map(s=>({path:s.path,headerRow:s.name==='Сотрудники'?3:1,
            columns:Array.from({length:s.maxCol},(_,i)=>({mode:i<2?mode:'keep',category:i===0?'Сотрудник':'Скрыто'}))}));
          const result=await S.Engine.anonymize(book,settings,'source.xlsx',true);
          const restored=mode==='alias'?await S.Engine.restore(await result.blob.arrayBuffer(),result.key):null;
          async function b64(blob){return await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}
          return {anon:await b64(result.blob),key:result.key,
            restored:restored?await b64(restored.blob):null, count:result.replaced};
        }''', {'data': source64, 'mode': mode})
        anon = self.work / 'anon.xlsx'
        anon.write_bytes(base64.b64decode(value['anon']))
        if value['restored']:
            (self.work / 'restored.xlsx').write_bytes(base64.b64decode(value['restored']))
        return value, anon

    def test_builtin_selftests(self):
        results = self.page.evaluate('async()=>await SafeCycle.SelfTest.run()')
        self.assertEqual(len(results), 7)
        self.assertTrue(all(result['pass'] for result in results), results)

    def test_array_formula_output_column_is_protected(self):
        path = self.work / 'array.xlsx'
        wb = xlsxwriter.Workbook(path)
        ws = wb.add_worksheet('Массив')
        ws.write_row('A1', ['Расчёт', 'Результат'])
        ws.write_array_formula('A2:B3', '={1,2;3,4}')
        wb.close()
        data = base64.b64encode(path.read_bytes()).decode()
        for mode in ['alias', 'delete']:
            response = self.page.evaluate('''async ({data,mode})=>{
              const S=SafeCycle,book=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);
              const before=await book.archive.text(book.sheets[0].path);
              try{
                await S.Engine.anonymize(book,[{path:book.sheets[0].path,headerRow:1,
                  columns:[{mode:'keep',category:'Скрыто'},{mode,category:'Скрыто'}]}],'array.xlsx');
                return {error:null};
              }catch(e){return {error:e.message,unchanged:before===await book.archive.text(book.sheets[0].path)}}
            }''', {'data': data, 'mode': mode})
            self.assertIsNotNone(response['error'], 'Разрешено менять ячейку результата формулы массива')
            self.assertIn('формул', response['error'])
            self.assertTrue(response['unchanged'])

    def test_category_control_matches_column_mode(self):
        self.page.locator('#demo').click()
        self.page.locator('#config-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        mode = self.page.get_by_label('Режим Сотрудники D', exact=True)
        category = self.page.get_by_label('Категория Сотрудники D', exact=True)
        self.assertEqual(mode.input_value(), 'keep')
        self.assertTrue(category.is_disabled(), 'Категория должна быть недоступна для «Оставить»')
        mode.select_option('alias')
        self.assertTrue(category.is_enabled())
        mode.select_option('delete')
        self.assertTrue(category.is_disabled())
        self.page.get_by_text('Для службы ИБ · ограничения и самопроверка', exact=True).click()
        self.page.locator('#selftest').click()
        self.page.locator('#busy').wait_for(state='hidden')
        self.assertTrue(category.is_disabled(), 'Самопроверка не должна менять состояние настройки')

    def test_cancel_and_retry_does_not_publish_partial_result(self):
        self.page.locator('#demo').click()
        self.page.locator('#config-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.evaluate('''()=>{
          document.getElementById('process').click();
          document.getElementById('cancel').click();
        }''')
        self.page.locator('#busy').wait_for(state='hidden')
        self.assertIn('отменена', self.page.locator('#status').inner_text())
        self.assertTrue(self.page.locator('#result-card').is_hidden())
        self.assertTrue(self.page.locator('#download-file').is_disabled())
        self.assertTrue(self.page.locator('#process').is_enabled())
        self.page.locator('#process').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertIn('Заменено ячеек: 12', self.page.locator('#report').inner_text())

    def test_exact_roundtrip_and_shared_string_variants(self):
        result, anon = self.output()
        self.assertEqual(result['count'], 8)
        anonymous = openpyxl.load_workbook(anon)
        self.assertEqual(anonymous['Сотрудники']['A4'].value, anonymous['Сотрудники']['A5'].value)
        self.assertEqual(anonymous['Сотрудники']['A4'].value, anonymous['Скрыто']['A2'].value)
        self.assertEqual(anonymous['Сотрудники']['C7'].value, '=SUM(C4:C6)')
        original = openpyxl.load_workbook(self.source)
        restored = openpyxl.load_workbook(self.work / 'restored.xlsx')
        for ws in original:
            back = restored[ws.title]
            self.assertEqual(ws.sheet_state, back.sheet_state)
            for row in ws:
                for cell in row:
                    returned = back[cell.coordinate]
                    self.assertEqual(cell.value, returned.value, f'{ws.title}!{cell.coordinate}')
                    self.assertEqual(cell.data_type, returned.data_type)
                    self.assertEqual(cell.style_id, returned.style_id)
                    self.assertEqual(cell.number_format, returned.number_format)
        self.assertEqual(restored['Сотрудники']['B4'].value, '00123')
        self.assertEqual(restored['Сотрудники']['B5'].value, '00007')

    def test_retains_unmodified_parts_and_structure(self):
        _, anon = self.output()
        before, after = parts(self.source), parts(anon)
        changed = {'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml',
                   'xl/sharedStrings.xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml'}
        self.assertEqual(before.keys(), after.keys())
        for name in before.keys() - changed:
            if name.endswith('.xml') or name.endswith('.rels') or name.endswith('.vml'):
                continue  # stage 4 serializes and inspects XML; assertions below cover retained structure.
            self.assertEqual(before[name], after[name], name)
        wb = openpyxl.load_workbook(anon)
        ws = wb['Сотрудники']
        self.assertEqual(ws.freeze_panes, 'B4')
        self.assertTrue(ws.row_dimensions[5].hidden)
        self.assertTrue(ws.column_dimensions['B'].hidden)
        self.assertEqual(ws.column_dimensions['A'].width,
                         openpyxl.load_workbook(self.source)['Сотрудники'].column_dimensions['A'].width)
        self.assertEqual(str(ws.merged_cells), 'A1:F1')
        self.assertEqual(len(ws.conditional_formatting), 1)
        self.assertEqual(len(ws.data_validations.dataValidation), 1)
        self.assertEqual(len(ws._charts), 1)
        self.assertTrue(ws['A4'].comment.author.startswith('Сотрудник_'))

    def test_unused_shared_strings_are_cleared(self):
        result, anon = self.output()
        output = parts(anon)
        # В этом этапе диаграммы и комментарии намеренно не проверяются на утечки.
        protected = ['xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml']
        for entry in result['key']['entries']:
            for original in entry['variants']:
                for name in protected:
                    self.assertNotIn(original, output[name].decode(), (name, original))

    def test_delete_is_not_written_to_key(self):
        result, anon = self.output('delete')
        wb = openpyxl.load_workbook(anon)
        self.assertIsNone(wb['Сотрудники']['A4'].value)
        self.assertIsNone(wb['Сотрудники']['B4'].value)
        self.assertEqual(wb['Сотрудники']['C4'].value, 80000)
        self.assertEqual([e['original'] for e in result['key']['entries']], ['Автор теста'])
        self.assertEqual(result['key']['cells'], [])
        self.assertNotIn('00123', json.dumps(result['key'], ensure_ascii=False))

    def test_restoration_finds_moved_cells_and_labels_in_text(self):
        result, anon = self.output()
        wb = openpyxl.load_workbook(anon)
        label = wb['Сотрудники']['A4'].value
        wb['Сотрудники']['A4'] = None
        wb['Сотрудники']['J20'] = label
        wb['Сотрудники']['K21'] = f'Ответ ИИ: [{label}]'
        edited = self.work / 'edited.xlsx'
        wb.save(edited)
        payload = self.page.evaluate('''async ({data,key})=>{
          const buffer=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer;
          const out=await SafeCycle.Engine.restore(buffer,key);
          const back=await SafeCycle.Xlsx.load(await out.blob.arrayBuffer());
          const s=back.sheets[0];return {same:out.sameHash,
          moved:SafeCycle.Xlsx.value(s.byRow.get(20).get(9),back.strings),
          text:SafeCycle.Xlsx.value(s.byRow.get(21).get(10),back.strings)}
        }''', {'data': base64.b64encode(edited.read_bytes()).decode(), 'key': result['key']})
        self.assertFalse(payload['same'])
        self.assertEqual(payload['moved'], ' Ёлкин Иван Иванович ')
        self.assertEqual(payload['text'], 'Ответ ИИ:  Ёлкин Иван Иванович ')

    def test_corrupted_checksum_is_rejected(self):
        raw = bytearray(self.source.read_bytes())
        with zipfile.ZipFile(self.source) as archive:
            info = archive.getinfo('xl/sharedStrings.xml')
            pos = info.header_offset
            name_len = int.from_bytes(raw[pos + 26:pos + 28], 'little')
            extra_len = int.from_bytes(raw[pos + 28:pos + 30], 'little')
            raw[pos + 30 + name_len + extra_len + info.compress_size // 2] ^= 4
        error = self.page.evaluate('''async data=>{try{
          await SafeCycle.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);
          return null;
        }catch(e){return e.message}}''', base64.b64encode(raw).decode())
        self.assertIsNotNone(error)
        self.assertRegex(error, 'поврежд|Поврежд|Контрольная')

    def test_existing_label_like_text_does_not_collide(self):
        response = self.page.evaluate('''async()=>{
          const S=SafeCycle,original=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer());
          const cell=original.sheets[0].byRow.get(2).get(2);
          S.Xlsx.setText(cell,'Сотрудник_0001');
          original.archive.set(original.sheets[0].path,S.Xml.serialize(original.sheets[0].doc));
          const book=await S.Xlsx.load(await (await original.archive.write()).arrayBuffer());
          const settings=book.sheets.map(s=>({path:s.path,headerRow:1,
            columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'alias':'keep',category:'Сотрудник'}))}));
          const out=await S.Engine.anonymize(book,settings,'collision.xlsx',false);
          const back=await S.Xlsx.load(await (await S.Engine.restore(await out.blob.arrayBuffer(),out.key)).blob.arrayBuffer());
          return {labels:out.key.entries.map(e=>e.label),unchanged:S.Xlsx.value(back.sheets[0].byRow.get(2).get(2),back.strings)};
        }''')
        self.assertNotIn('Сотрудник_0001', response['labels'])
        self.assertEqual(response['unchanged'], 'Сотрудник_0001')

    def test_zip64_small_entry_is_readable(self):
        packed = io.BytesIO()
        with zipfile.ZipFile(packed, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            with archive.open('test.xml', 'w', force_zip64=True) as item:
                item.write('Тест чтения'.encode())
        value = self.page.evaluate('''async data=>{
          const a=await SafeCycle.Zip.Archive.read(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer);
          return a.text('test.xml');
        }''', base64.b64encode(packed.getvalue()).decode())
        self.assertEqual(value, 'Тест чтения')

    def test_interface_and_downloads(self):
        self.page.locator('#file').set_input_files(self.source)
        self.page.locator('#config-card').wait_for(state='visible')
        self.page.locator('#sheets input[type=number]').first.fill('3')
        self.page.locator('#sheets input[type=number]').first.dispatch_event('change')
        self.assertEqual(self.page.locator('img').count(), 0, 'Текст ячейки превратился в разметку')
        self.page.locator('#process').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.assertTrue(self.page.locator('#download-file').is_disabled())
        self.page.locator('#ack').check()
        with self.page.expect_download() as download:
            self.page.locator('#download-file').click()
        output = self.work / 'ui.xlsx'
        download.value.save_as(output)
        self.assertEqual(download.value.suggested_filename, 'source_anon.xlsx')
        openpyxl.load_workbook(output)
        self.page.locator('#key-encryption').uncheck()
        with self.page.expect_download() as download:
            self.page.locator('#download-key').click()
        key_file = self.work / 'ui-key.json'
        download.value.save_as(key_file)
        self.page.locator('#restore-tab').click()
        self.page.locator('#key-file').set_input_files(key_file)
        self.page.locator('#return-file').set_input_files(output)
        self.page.locator('#restore').click()
        self.page.locator('#download-restored').wait_for(state='visible')
        self.assertIn('Отпечаток совпал', self.page.locator('#restore-report').inner_text())
        with self.page.expect_download() as download:
            self.page.locator('#download-restored').click()
        back = self.work / 'ui-restored.xlsx'
        download.value.save_as(back)
        self.assertEqual(openpyxl.load_workbook(back)['Сотрудники']['A4'].value, ' Ёлкин Иван Иванович ')
        self.page.locator('#anon-tab').click()
        self.page.locator('#reset').click()
        self.assertTrue(self.page.locator('#config-card').is_hidden())
        self.assertEqual(self.page.locator('#key-file').input_value(), '')
        self.assertEqual(self.page.locator('#return-file').input_value(), '')

    def test_password_container_and_wrong_key_errors(self):
        error = self.page.evaluate('''async()=>{try{
          await SafeCycle.Xlsx.load(new Uint8Array([208,207,17,224,161,177,26,225]).buffer);
          return null;
        }catch(e){return e.message}}''')
        self.assertIn('паролем', error)
        error = self.page.evaluate('''()=>{try{SafeCycle.Key.validate({schema:'bad'});return null}catch(e){return e.message}}''')
        self.assertIn('Ключ', error)

    def test_csp_and_no_persistent_storage(self):
        source = HTML.read_text()
        self.assertIn("connect-src 'none'", source)
        self.assertIn("default-src 'none'", source)
        self.assertNotRegex(source, r'localStorage|sessionStorage|indexedDB|document\.cookie|<script[^>]+src=|<link[^>]+href=')
        self.assertEqual(self.page.evaluate('performance.getEntriesByType("resource").length'), 0)
        # Отдельная страница: сообщения CSP от намеренного запроса ожидаются.
        security_page = self.browser.new_page()
        attempted = []
        security_page.route('https://example.invalid/**', lambda route: (attempted.append(route.request.url), route.abort()))
        security_page.goto(self.url)
        blocked = security_page.evaluate('''async()=>{try{await fetch('https://example.invalid/check');return false}catch{return true}}''')
        self.assertTrue(blocked)
        self.assertEqual(attempted, [], 'Сетевая операция дошла до отправки несмотря на CSP')
        security_page.close()

    def test_generated_demo_opens_independently(self):
        data = self.page.evaluate('''async()=>{
          const blob=await SafeCycle.Demo.create();
          return await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)});
        }''')
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data)))
        self.assertEqual(wb['Сотрудники']['B2'].value, '00001')
        self.assertEqual(wb['Сотрудники']['D7'].value, '=SUM(D2:D6)')
        self.assertEqual(wb['Скрытый лист'].sheet_state, 'hidden')

    def test_layout_at_tablet_width(self):
        self.page.set_viewport_size({'width': 800, 'height': 1000})
        self.page.locator('#demo').click()
        self.page.locator('#config-card').wait_for(state='visible')
        self.assertFalse(self.page.evaluate('document.documentElement.scrollWidth>innerWidth'))
        self.page.screenshot(path=str(self.work / 'tablet.png'), full_page=True)

    def test_libreoffice_conversion_when_available(self):
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if not executable:
            self.skipTest('LibreOffice отсутствует; совместимость с ним не подтверждена')
        _, anon = self.output()
        destination = self.work / 'lo'
        destination.mkdir(exist_ok=True)
        run = subprocess.run([executable, f'-env:UserInstallation={self.work.as_uri()}/lo-profile',
                              '--headless', '--convert-to', 'xlsx', '--outdir', str(destination),
                              str(anon), str(self.work / 'restored.xlsx')],
                             capture_output=True, text=True, timeout=60)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue((destination / anon.name).exists(), run.stdout + run.stderr)
        openpyxl.load_workbook(destination / anon.name)
        self.assertTrue((destination / 'restored.xlsx').exists(), run.stdout + run.stderr)
        self.assertEqual(openpyxl.load_workbook(destination / 'restored.xlsx')['Сотрудники']['A4'].value,
                         ' Ёлкин Иван Иванович ')


if __name__ == '__main__':
    unittest.main(verbosity=2)
