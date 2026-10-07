"""Скрытые части XLSX и запрет скачивания; только вымышленные данные."""
import base64
import io
import json
from pathlib import Path
import shutil
import subprocess
import unittest
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import xlsxwriter
from playwright.sync_api import expect
import test_stage1 as stage1

S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
NAME = 'Иванов Иван Иванович'
EMAIL = 'ivan@example.test'


def fixture(path):
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet('Данные')
    ws.write_row('A1', ['ФИО', 'Email', 'СеверПроект', 'Формула', 'Комментарий'])
    ws.write_row('A2', [NAME, EMAIL, 175123, '', 'Передать код 876543210987'])
    ws.write_row('A3', ['Петрова Анна Сергеевна', 'anna@example.test', 201123, '', 'Работает в Звездограде'])
    ws.write_formula('D2', '=A2', None, NAME)
    ws.write_formula('D3', '=C2+C3', None, 376246)
    ws.write_comment('A2', 'Согласовано с Ивановым Иваном Ивановичем: '+EMAIL, {'author': NAME})
    ws.set_header('&L'+NAME+'&R'+EMAIL)
    ws.set_footer('&LСеверПроект')
    ws.data_validation('E2:E3', {'validate': 'list', 'source': [NAME, EMAIL]})
    ws.write_url('F2', 'https://private.example.test/'+EMAIL, string='Ссылка')
    ws.set_row(2, None, None, {'hidden': True})
    ws.set_column('B:B', 20, None, {'hidden': True})
    ws.add_table('A1:E3', {'name': 'СеверПроект', 'columns': [{'header': h} for h in ['ФИО', 'Email', 'СеверПроект', 'Формула', 'Комментарий']]})
    chart = wb.add_chart({'type': 'column'})
    chart.add_series({'values': '=Данные!$C$2:$C$3', 'categories': '=Данные!$A$2:$A$3', 'name': NAME})
    chart.set_title({'name': 'Отчет '+NAME})
    ws.insert_chart('H2', chart)
    # Small real PNG: removed as an unexamined binary image.
    png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aN1sAAAAASUVORK5CYII=')
    ws.insert_image('H20', 'portrait.png', {'image_data': io.BytesIO(png), 'description': NAME})
    wb.define_name('Имя', '="'+NAME+'"')
    wb.define_name('СеверПроект_Имя', '=Данные!$A$2')
    wb.set_properties({'author': NAME, 'company': 'СеверПроект', 'manager': NAME})
    wb.set_custom_property('Email', EMAIL)
    hidden = wb.add_worksheet('Скрыто')
    hidden.hide()
    hidden.write_row('A1', ['ФИО', 'Email'])
    hidden.write_rich_string('A2', wb.add_format({'bold': True}), 'Иванов ', 'Иван Иванович')
    hidden.write('B2', EMAIL)
    very = wb.add_worksheet('Полностью скрыто')
    very.very_hidden()
    very.write_row('A1', ['ФИО', 'Email'])
    very.write_row('A2', [NAME, EMAIL])
    wb.close()


def rewrite(path, changes):
    data = stage1.parts(path)
    data.update(changes)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, payload in data.items():
            z.writestr(name, payload)


def text(data):
    return data.decode('utf-8')


class Stage4(unittest.TestCase):
    setUpClass = classmethod(stage1.Stage1.setUpClass.__func__)
    tearDownClass = classmethod(stage1.Stage1.tearDownClass.__func__)
    setUp = stage1.Stage1.setUp
    tearDown = stage1.Stage1.tearDown

    def process(self, source=None, mode='keep', restore=True):
        source = source or self.work / 'hidden.xlsx'
        if not source.exists():
            fixture(source)
        payload = self.page.evaluate('''async ({data,mode,restore})=>{
          const S=SafeCycle,b=Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,book=await S.Xlsx.load(b);
          const settings=book.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:c===0?'alias':c===1?'alias':c===2?mode:'keep',category:c===0?'Сотрудник':c===1?'Email':'Скрыто',rangeStep:50000}))}));
          const out=await S.Engine.anonymize(book,settings,'hidden.xlsx',false,undefined,undefined,{dictionary:[{value:'СеверПроект',category:'Проект'}]});
          const back=restore?await S.Engine.restore(await out.blob.arrayBuffer(),JSON.parse(JSON.stringify(out.key))):null;
          async function b64(blob){return new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(blob)})}
          return {anon:await b64(out.blob),back:back?await b64(back.blob):null,key:out.key,audit:out.audit,hidden:out.hiddenChanges};
        }''', {'data': base64.b64encode(source.read_bytes()).decode(), 'mode': mode, 'restore': restore})
        anon = self.work / 'hidden-anon.xlsx'
        anon.write_bytes(base64.b64decode(payload['anon']))
        if payload['back']:
            (self.work / 'hidden-back.xlsx').write_bytes(base64.b64decode(payload['back']))
        return payload, anon

    def test_comments_headers_lists_charts_props_links_hidden_sheets(self):
        payload, anon = self.process()
        self.assertFalse(payload['audit']['blocked'], payload['audit']['findings'])
        xml = stage1.parts(anon)
        for part, data in xml.items():
            if part.endswith(('.xml', '.rels', '.vml')):
                for secret in [NAME, EMAIL, 'Ивановым Иваном Ивановичем', 'СеверПроект']:
                    self.assertNotIn(secret, text(data), part)
        self.assertNotIn('docProps/custom.xml', xml)
        self.assertFalse(any(p.startswith('xl/media/') for p in xml))
        self.assertNotIn('numCache', text(xml['xl/charts/chart1.xml']))
        wb = openpyxl.load_workbook(anon)
        ws = wb['Данные']
        label = ws['A2'].value
        self.assertEqual(ws['A2'].comment.author, label)
        self.assertIn('['+label+']', ws['A2'].comment.text)
        self.assertIn('['+label+']', ws.oddHeader.left.text)
        self.assertIn('['+label+']', ws.data_validations.dataValidation[0].formula1)
        self.assertEqual(ws['A2'].value, wb['Скрыто']['A2'].value)
        self.assertEqual(ws['A2'].value, wb['Полностью скрыто']['A2'].value)
        self.assertEqual(wb['Скрыто'].sheet_state, 'hidden')
        self.assertEqual(wb['Полностью скрыто'].sheet_state, 'veryHidden')
        self.assertTrue(ws.row_dimensions[3].hidden)
        self.assertTrue(ws.column_dimensions['B'].hidden)
        self.assertIsNone(ws['F2'].hyperlink)
        self.assertEqual(ws['D2'].value, '=A2')
        self.assertIsNone(openpyxl.load_workbook(anon, data_only=True)['Данные']['D2'].value)
        table = list(ws.tables.values())[0]
        self.assertEqual(table.tableColumns[2].name, ws['C1'].value)
        self.assertTrue(table.displayName.startswith('Проект_'))
        self.assertTrue(wb.calculation.forceFullCalc)
        self.assertIsNone(wb.properties.creator)

    def test_exact_hidden_roundtrip_uses_original_case_and_retains_formula(self):
        payload, _ = self.process()
        wb = openpyxl.load_workbook(self.work / 'hidden-back.xlsx')
        ws = wb['Данные']
        self.assertEqual(ws['A2'].comment.author, NAME)
        self.assertIn('Ивановым Иваном Ивановичем', ws['A2'].comment.text)
        self.assertEqual(ws['C1'].value, 'СеверПроект')
        self.assertEqual(list(ws.tables.values())[0].tableColumns[2].name, 'СеверПроект')
        self.assertIn(NAME, ws.oddHeader.left.text)
        self.assertIn(NAME, ws.data_validations.dataValidation[0].formula1)
        self.assertEqual(openpyxl.load_workbook(self.work / 'hidden-back.xlsx', data_only=True)['Данные']['D2'].value, NAME)
        self.assertEqual(ws['D2'].value, '=A2')
        self.assertIsNone(wb.properties.creator)
        self.assertIsNone(ws['F2'].hyperlink)

    def test_range_values_in_metadata_and_formula_caches_are_not_saved(self):
        source = self.work / 'range-hidden.xlsx'
        fixture(source)
        data = stage1.parts(source)
        comments = text(data['xl/comments1.xml']).replace('</text>', '<t> 175123</t></text>')
        rewrite(source, {'xl/comments1.xml': comments})
        payload, anon = self.process(source, mode='range')
        self.assertNotIn('175123', json.dumps(payload['key'], ensure_ascii=False))
        self.assertEqual(payload['key']['formulaCaches'], [])
        for part, data in stage1.parts(anon).items():
            if part.endswith(('.xml', '.rels', '.vml')):
                self.assertNotIn('175123', text(data), part)
        self.assertFalse(payload['audit']['blocked'], payload['audit']['findings'])

    def test_unknown_xml_and_case_normalization_report_part(self):
        source = self.work / 'extension.xlsx'
        fixture(source)
        rewrite(source, {'xl/unknown.xml': '<extra><secret>ивановым   иваном ивановичем</secret><secret>'+EMAIL+'</secret></extra>'})
        payload, _ = self.process(source)
        self.assertTrue(payload['audit']['blocked'])
        found = [f for f in payload['audit']['findings'] if f['part'] == 'xl/unknown.xml']
        self.assertTrue(any('ивановым' in f['match'] for f in found), found)
        self.assertTrue(any(EMAIL == f['match'] for f in found), found)

    def test_xml_comments_and_encoded_attributes_are_checked(self):
        source = self.work / 'comments-xml.xlsx'
        fixture(source)
        rewrite(source, {'xl/unknown.xml': '<extra value="ivan&#64;example.test"><!--'+NAME+'--></extra>'})
        payload, _ = self.process(source)
        self.assertTrue(any(f['part'] == 'xl/unknown.xml' and f['match'] == NAME for f in payload['audit']['findings']))
        self.assertTrue(any(f['part'] == 'xl/unknown.xml' and f['match'] == EMAIL for f in payload['audit']['findings']))

    def test_private_formula_and_external_function_block_download(self):
        source = self.work / 'formula-private.xlsx'
        fixture(source)
        data = stage1.parts(source)
        sheet = text(data['xl/worksheets/sheet1.xml']).replace('<f>A2</f>', '<f>HYPERLINK("https://example.test/'+EMAIL+'","'+NAME+'")</f>')
        rewrite(source, {'xl/worksheets/sheet1.xml': sheet})
        payload, _ = self.process(source)
        self.assertTrue(payload['audit']['blocked'])
        self.assertTrue(any(f['cell'] == 'D2' and f['kind'] == 'Внешняя формула' for f in payload['audit']['findings'] if 'cell' in f))
        self.assertTrue(any(f.get('cell') == 'D2' and f['match'] == EMAIL for f in payload['audit']['findings']))

    def test_opaque_part_blocks_and_source_is_unchanged(self):
        source = self.work / 'opaque.xlsx'
        fixture(source)
        rewrite(source, {'xl/unknown.bin': b'secret payload'})
        before = source.read_bytes()
        payload, _ = self.process(source)
        self.assertEqual(before, source.read_bytes())
        self.assertTrue(any(f['part'] == 'xl/unknown.bin' and f['kind'] == 'Непроверенная часть' for f in payload['audit']['findings']))

    def test_human_review_kept_numeric_cell_and_capital_word(self):
        payload, _ = self.process()
        values = [c['value'] for c in payload['audit']['review']]
        self.assertIn('876543210987', values)
        self.assertIn('Звездограде', values)
        self.assertNotIn('Сотрудник', values)
        result = self.page.evaluate('async()=>{const s=SafeCycle;return await s.Engine.anonymizeText("Обсудить с Звездоградом и кодом 123456789987",{dictionary:[{category:"Скрыто",value:"Звездоградом",exact:true}]},false)}')
        self.assertIn('123456789987', [r['value'] for r in result['audit']['review']])

    def test_ui_leak_guard_requires_explicit_override_and_resets(self):
        source = self.work / 'guard.xlsx'
        fixture(source)
        rewrite(source, {'xl/unknown.xml': '<extra>'+EMAIL+'</extra>'})
        self.page.locator('#file').set_input_files(source)
        self.page.locator('#config-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.get_by_label('Режим Данные A', exact=True).select_option('alias')
        self.page.get_by_label('Категория Данные A', exact=True).select_option('Сотрудник')
        self.page.locator('#process').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#ack').check()
        expect(self.page.locator('#download-file')).to_be_disabled()
        expect(self.page.locator('#download-key')).to_be_disabled()
        self.assertIn('xl/unknown.xml', self.page.locator('#leak-list').inner_text())
        self.page.locator('#download-override').check()
        expect(self.page.locator('#download-file')).to_be_enabled()
        with self.page.expect_download() as download:
            self.page.locator('#download-file').click()
        self.assertEqual(download.value.suggested_filename, 'guard_anon.xlsx')
        self.page.locator('#process').click()
        self.page.locator('#busy').wait_for(state='hidden')
        self.assertFalse(self.page.locator('#download-override').is_checked())
        expect(self.page.locator('#download-file')).to_be_disabled()

    def test_ui_one_click_hide_rebuilds_source_and_reset_clears_report(self):
        self.page.locator('#source-text').click()
        self.page.locator('#raw-text').fill('Написать ivan@example.test и передать в Звездоград')
        self.page.locator('#process-text').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.get_by_role('button', name='Скрыть Звездоград', exact=True).click()
        self.page.locator('#busy').wait_for(state='hidden')
        self.assertNotIn('Звездоград', self.page.locator('#anon-text').input_value())
        self.assertIn('Скрыто_', self.page.locator('#anon-text').input_value())
        self.page.locator('#reset').click()
        self.assertEqual(self.page.locator('#manual-review').inner_text(), '')
        self.assertEqual(self.page.locator('#leak-list').inner_text(), '')

    def test_outputs_open_and_recalculate_in_libreoffice(self):
        executable = shutil.which('soffice') or shutil.which('libreoffice')
        self.assertIsNotNone(executable)
        _, anon = self.process()
        destination = self.work / 'stage4-lo'
        destination.mkdir(exist_ok=True)
        files = [anon, self.work / 'hidden-back.xlsx']
        run = subprocess.run([executable, f'-env:UserInstallation={self.work.as_uri()}/stage4-profile', '--headless', '--convert-to', 'xlsx', '--outdir', str(destination), *map(str, files)], capture_output=True, text=True, timeout=60)
        self.assertEqual(run.returncode, 0, run.stdout+run.stderr)
        for source in files:
            converted = destination / source.name
            self.assertTrue(converted.exists(), run.stdout+run.stderr)
            wb = openpyxl.load_workbook(converted)
            self.assertEqual(wb['Данные']['D2'].value, '=A2')
            self.assertEqual(len(wb['Данные']._charts), 1)
        cached = openpyxl.load_workbook(destination / anon.name, data_only=True)
        self.assertTrue(cached['Данные']['D2'].value.startswith('Сотрудник_'))

    def test_threaded_comments_person_ids_and_editable_authors(self):
        source = self.work / 'threaded.xlsx'
        fixture(source)
        ns = 'http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments'
        old_id = '{11111111-2222-3333-4444-555555555555}'
        person = f'<personList xmlns="{ns}"><person displayName="{NAME}" id="{old_id}" userId="{EMAIL}" providerId="AD"/></personList>'
        comment = f'<ThreadedComments xmlns="{ns}"><threadedComment ref="A2" personId="{old_id}" dT="2026-10-07T12:00:00Z" id="{{22222222-2222-3333-4444-555555555555}}"><text>Ответ {NAME}, {EMAIL}</text></threadedComment></ThreadedComments>'
        data = stage1.parts(source)
        rel = text(data['xl/_rels/workbook.xml.rels']).replace('</Relationships>', '<Relationship Id="persons" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" Target="persons/person.xml"/></Relationships>')
        srel = text(data['xl/worksheets/_rels/sheet1.xml.rels']).replace('</Relationships>', '<Relationship Id="threaded" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" Target="../threadedComments/comment.xml"/></Relationships>')
        types = text(data['[Content_Types].xml']).replace('</Types>', '<Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/><Override PartName="/xl/threadedComments/comment.xml" ContentType="application/vnd.ms-excel.threadedcomments+xml"/></Types>')
        rewrite(source, {'xl/persons/person.xml': person, 'xl/threadedComments/comment.xml': comment, 'xl/_rels/workbook.xml.rels': rel, 'xl/worksheets/_rels/sheet1.xml.rels': srel, '[Content_Types].xml': types})
        payload, anon = self.process(source)
        self.assertFalse(payload['audit']['blocked'], payload['audit']['findings'])
        data = stage1.parts(anon)
        p = ET.fromstring(data['xl/persons/person.xml'])[0]
        c = ET.fromstring(data['xl/threadedComments/comment.xml'])[0]
        self.assertNotEqual(p.attrib['id'], old_id)
        self.assertEqual(p.attrib['id'], c.attrib['personId'])
        self.assertEqual(p.attrib['userId'], '')
        self.assertEqual(p.attrib['displayName'], openpyxl.load_workbook(anon)['Данные']['A2'].value)
        back = stage1.parts(self.work / 'hidden-back.xlsx')
        self.assertEqual(ET.fromstring(back['xl/persons/person.xml'])[0].attrib['displayName'], NAME)
        self.assertIn(NAME, text(back['xl/threadedComments/comment.xml']))

    def test_external_link_path_neutralized_cache_removed_indices_preserved(self):
        source = self.work / 'external.xlsx'
        fixture(source)
        data = stage1.parts(source)
        ext = f'<externalLink xmlns="{S}" xmlns:r="{R}"><externalBook r:id="external"><sheetNames><sheetName val="Лист"/></sheetNames><sheetDataSet><sheetData sheetId="0"><row r="1"><cell r="A1" t="str"><v>{NAME}</v></cell></row></sheetData></sheetDataSet></externalBook></externalLink>'
        rel = f'<Relationships xmlns="{P}"><Relationship Id="external" Type="{R}/externalLinkPath" Target="file:///private/{EMAIL}/book.xlsx" TargetMode="External"/></Relationships>'
        w = text(data['xl/workbook.xml']).replace('<calcPr', '<externalReferences><externalReference r:id="extLink"/></externalReferences><calcPr')
        wrel = text(data['xl/_rels/workbook.xml.rels']).replace('</Relationships>', f'<Relationship Id="extLink" Type="{R}/externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>')
        types = text(data['[Content_Types].xml']).replace('</Types>', '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>')
        rewrite(source, {'xl/externalLinks/externalLink1.xml': ext, 'xl/externalLinks/_rels/externalLink1.xml.rels': rel, 'xl/workbook.xml': w, 'xl/_rels/workbook.xml.rels': wrel, '[Content_Types].xml': types})
        payload, anon = self.process(source)
        self.assertFalse(payload['audit']['blocked'], payload['audit']['findings'])
        data = stage1.parts(anon)
        self.assertNotIn(EMAIL, text(data['xl/externalLinks/_rels/externalLink1.xml.rels']))
        self.assertIn('file:///removed-external-source.xlsx', text(data['xl/externalLinks/_rels/externalLink1.xml.rels']))
        self.assertNotIn('sheetDataSet', text(data['xl/externalLinks/externalLink1.xml']))
        wb = openpyxl.load_workbook(anon)
        self.assertEqual(len(wb._external_links), 1)
        self.assertEqual(wb._external_links[0].file_link.Target, 'file:///removed-external-source.xlsx')

    def test_pivot_cache_blank_slots_refresh_and_independent_readers(self):
        from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
        from openpyxl.pivot.fields import Text, Number, Index
        from openpyxl.pivot.record import RecordList, Record
        from openpyxl.pivot.table import TableDefinition, Location, PivotField, FieldItem, RowColField, DataField
        source = self.work / 'pivot.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Данные'
        ws.append(['ФИО', 'Оклад'])
        ws.append([NAME, 175123])
        ws.append(['Петрова Анна Сергеевна', 201123])
        cache = CacheDefinition(cacheSource=CacheSource(type='worksheet', worksheetSource=WorksheetSource(ref='A1:B3', sheet='Данные')), cacheFields=[CacheField(name='ФИО', sharedItems=SharedItems(_fields=[Text(v=NAME), Text(v='Петрова Анна Сергеевна')], containsString=True)), CacheField(name='Оклад', sharedItems=SharedItems(_fields=[Number(v=175123), Number(v=201123)], containsNumber=True))], recordCount=2, refreshedBy=NAME)
        cache.records = RecordList(r=[Record(_fields=[Index(v=0), Number(v=175123)]), Record(_fields=[Index(v=1), Number(v=201123)])])
        pivot = TableDefinition(name='Сводная', cacheId=1, dataCaption='Сумма', location=Location(ref='D1:E4', firstHeaderRow=1, firstDataRow=1, firstDataCol=1), pivotFields=[PivotField(axis='axisRow', items=[FieldItem(x=0), FieldItem(x=1)]), PivotField(dataField=True)], rowFields=[RowColField(x=0)], dataFields=[DataField(name='Сумма', fld=1)])
        pivot.cache = cache
        ws.add_pivot(pivot)
        wb.save(source)
        payload, anon = self.process(source, restore=True)
        self.assertFalse(payload['audit']['blocked'], payload['audit']['findings'])
        data = stage1.parts(anon)
        definition = next(v for p, v in data.items() if 'pivotCacheDefinition' in p)
        records = next(v for p, v in data.items() if 'pivotCacheRecords' in p)
        root = ET.fromstring(definition)
        self.assertEqual(root.attrib['refreshOnLoad'], '1')
        self.assertEqual(root.attrib['saveData'], '0')
        self.assertNotIn(NAME, text(definition))
        self.assertNotIn('175123', text(definition)+text(records))
        self.assertNotIn('201123', text(definition)+text(records))
        for items in root.findall(f'.//{{{S}}}sharedItems'):
            self.assertEqual(len(items), 2)
            self.assertTrue(all(n.tag == f'{{{S}}}m' for n in items))
        self.assertEqual(len(openpyxl.load_workbook(anon)['Данные']._pivots), 1)
        destination = self.work / 'pivot-lo'
        destination.mkdir(exist_ok=True)
        run = subprocess.run([shutil.which('soffice'), f'-env:UserInstallation={self.work.as_uri()}/pivot-profile', '--headless', '--convert-to', 'xlsx', '--outdir', str(destination), str(anon)], capture_output=True, text=True, timeout=60)
        self.assertEqual(run.returncode, 0, run.stdout+run.stderr)
        self.assertTrue((destination / anon.name).exists(), run.stdout+run.stderr)
        converted = openpyxl.load_workbook(destination / anon.name)
        self.assertEqual(len(converted['Данные']._pivots), 1)
        self.assertTrue(converted['Данные']['A2'].value.startswith('Сотрудник_'))

    def test_builtin_privacy_selfchecks_and_ui_demo(self):
        results = self.page.evaluate('async()=>await SafeCycle.PrivacySelfTest.run()')
        self.assertEqual(len(results), 4)
        self.assertTrue(all(r['pass'] for r in results), results)
        self.page.locator('#privacy-demo').click()
        self.page.locator('#config-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#process').click()
        self.page.locator('#result-card').wait_for(state='visible')
        self.page.locator('#busy').wait_for(state='hidden')
        self.page.locator('#hidden-details summary').click()
        self.assertIn('xl/comments1.xml', self.page.locator('#hidden-report').inner_text())
        self.assertNotIn('Скачивание заблокировано', self.page.locator('#audit-report').inner_text())
        self.page.locator('#ack').check()
        expect(self.page.locator('#download-file')).to_be_enabled()

    def test_numeric_payload_checked_without_confusing_structural_attributes(self):
        source = self.work / 'numeric-secret.xlsx'
        fixture(source)
        data = stage1.parts(source)
        # Keep numeric 0 unchanged by scaling; schema IDs also commonly equal zero.
        sheet = text(data['xl/worksheets/sheet1.xml']).replace('<v>175123</v>', '<v>0</v>')
        rewrite(source, {'xl/worksheets/sheet1.xml': sheet, 'xl/unknown.xml': '<extra value="175123"><secret>код 0</secret></extra>'})
        payload, _ = self.process(source, mode='scale')
        findings = payload['audit']['findings']
        self.assertTrue(any(f['part'] == 'xl/unknown.xml' and f['match'] == '0' for f in findings), findings)
        self.assertFalse(any(f['part'] == 'xl/styles.xml' for f in findings), findings)
        self.assertFalse(any(f['cell'] == 'C2' for f in findings if 'cell' in f), findings)

    def test_shared_formula_and_array_cached_results_all_cleared_and_restored(self):
        source = self.work / 'array-cache.xlsx'
        wb = xlsxwriter.Workbook(source)
        ws = wb.add_worksheet('Данные')
        ws.write_row('A1', ['ФИО', 'Расчёт', 'Расчёт'])
        ws.write('A2', NAME)
        ws.write_array_formula('B2:C3', '{={1,2;3,4}}', value=1)
        ws.write_number('C2', 2)
        ws.write_number('B3', 3)
        ws.write_number('C3', 4)
        wb.close()
        payload = self.page.evaluate('''async data=>{const S=SafeCycle,b=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),out=await S.Engine.anonymize(b,[{path:b.sheets[0].path,headerRow:1,columns:[{mode:'alias',category:'Сотрудник'},{mode:'keep'},{mode:'keep'}]}],'array.xlsx',false),anon=await S.Xlsx.load(await out.blob.arrayBuffer()),back=await S.Engine.restore(await out.blob.arrayBuffer(),out.key),book=await S.Xlsx.load(await back.blob.arrayBuffer());return {caches:out.key.stats.formulaCaches,values:['B2','C2','B3','C3'].map(ref=>S.Xlsx.value(anon.sheets[0].cells.find(c=>c.getAttribute('r')===ref),anon.strings)),back:['B2','C2','B3','C3'].map(ref=>S.Xlsx.value(book.sheets[0].cells.find(c=>c.getAttribute('r')===ref),book.strings)),formula:S.Xml.children(anon.sheets[0].byRow.get(2).get(1),'f')[0].textContent}}''', base64.b64encode(source.read_bytes()).decode())
        self.assertEqual(payload['caches'], 4)
        self.assertEqual(payload['values'], ['', '', '', ''])
        self.assertEqual(payload['back'], ['1', '2', '3', '4'])
        self.assertEqual(payload['formula'], '{1,2;3,4}')

    def test_deleted_person_cases_in_hidden_text_do_not_reappear_in_key(self):
        source = self.work / 'deleted-person.xlsx'
        fixture(source)
        result = self.page.evaluate('''async data=>{const S=SafeCycle,book=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),out=await S.Engine.anonymize(book,book.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'delete':'keep',category:'Скрыто'}))})),'deleted.xlsx',false);return {key:out.key,audit:out.audit,parts:await (async()=>{const a=await S.Zip.Archive.read(await out.blob.arrayBuffer());return await a.text('xl/comments1.xml')})()}}''', base64.b64encode(source.read_bytes()).decode())
        raw = json.dumps(result['key'], ensure_ascii=False)
        self.assertNotIn(NAME, raw)
        self.assertNotIn('Ивановым Иваном Ивановичем', raw)
        self.assertNotIn(NAME, result['parts'])
        self.assertNotIn('Ивановым Иваном Ивановичем', result['parts'])
        self.assertFalse(result['audit']['blocked'], result['audit']['findings'])

    def test_short_source_value_in_unknown_xml_text_is_not_skipped(self):
        result = self.page.evaluate('''async()=>{const S=SafeCycle,b=await S.Xlsx.load(await (await S.Demo.create()).arrayBuffer());b.archive.set('xl/unknown.xml','<extra>Отдел HR</extra>');const out=await S.Engine.anonymize(b,b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0||i===2?'alias':'keep',category:i===0?'Сотрудник':'Проект'}))})),'short.xlsx',false);return out.audit}''')
        self.assertTrue(result['blocked'])
        self.assertTrue(any(f['part'] == 'xl/unknown.xml' and f['match'] == 'HR' for f in result['findings']), result)

    def test_stage3_numeric_key_preserves_unmodified_metadata_labels(self):
        old = self.browser.new_page()
        try:
            old.goto(self.url.replace('v0.6.0', 'v0.3.0'))
            payload = old.evaluate('''async()=>{const S=SafeCycle,book=await S.Xlsx.load(await (await S.Demo.createNumbers()).arrayBuffer());book.archive.set('docProps/core.xml','<coreProperties xmlns="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>Сотрудник_0001</dc:creator></coreProperties>');const out=await S.Engine.anonymize(book,book.sheets.map((s,i)=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,c)=>({mode:i?'keep':['scale','keep','range','date','alias'][c],category:'Сотрудник',rangeStep:50000}))})),'stage3.xlsx',false,undefined,undefined,{crosscheck:false});return {key:out.key,data:await new Promise(resolve=>{const r=new FileReader();r.onload=()=>resolve(r.result.split(',')[1]);r.readAsDataURL(out.blob)})}}''')
        finally:
            old.close()
        self.assertEqual(payload['key']['version'], '0.3.0')
        result = self.page.evaluate('''async ({key,data})=>{const S=SafeCycle,out=await S.Engine.restore(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer,key),b=await S.Xlsx.load(await out.blob.arrayBuffer());return {sameHash:out.sameHash,money:S.Xlsx.value(b.sheets[0].byRow.get(2).get(0),b.strings),date:S.Xlsx.value(b.sheets[0].byRow.get(2).get(3),b.strings),code:S.Xlsx.value(b.sheets[0].byRow.get(2).get(4),b.strings),metadata:await b.archive.text('docProps/core.xml')}}''', payload)
        self.assertTrue(result['sameHash'])
        self.assertEqual(result['money'], '150000.25')
        self.assertEqual(result['date'], '45000')
        self.assertEqual(result['code'], '00123')
        self.assertIn('Сотрудник_0001', result['metadata'])
        self.assertNotIn('00123', result['metadata'])

    def test_deleted_data_in_phonetic_annotation_is_not_saved_in_key(self):
        source = self.work / 'phonetic-key.xlsx'
        fixture(source)
        data = stage1.parts(source)
        sst = text(data['xl/sharedStrings.xml']).replace('<si><t>'+EMAIL+'</t></si>', '<si><t>'+EMAIL+'</t><rPh sb="0" eb="4"><t>'+NAME+'</t></rPh></si>')
        self.assertNotEqual(sst, text(data['xl/sharedStrings.xml']))
        rewrite(source, {'xl/sharedStrings.xml': sst})
        result = self.page.evaluate('''async data=>{const S=SafeCycle,b=await S.Xlsx.load(Uint8Array.from(atob(data),c=>c.charCodeAt(0)).buffer),out=await S.Engine.anonymize(b,b.sheets.map(s=>({path:s.path,headerRow:1,columns:Array.from({length:s.maxCol},(_,i)=>({mode:i===0?'delete':i===1?'alias':'keep',category:'Email'}))})),'phonetic.xlsx',false);return {key:out.key,audit:out.audit}}''', base64.b64encode(source.read_bytes()).decode())
        raw = json.dumps(result['key'], ensure_ascii=False)
        self.assertNotIn(NAME, raw)
        self.assertNotIn('rPh', raw)
        self.assertFalse(result['audit']['blocked'], result['audit']['findings'])
